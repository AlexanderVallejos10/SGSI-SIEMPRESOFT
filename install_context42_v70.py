from pathlib import Path
import textwrap

ROOT = Path('.')

FILES = {
    'apps/context42/__init__.py': '',
    'apps/context42/apps.py': '''from django.apps import AppConfig


class Context42Config(AppConfig):
    default_auto_field = "django.db.models.BigAutoField"
    name = "apps.context42"
    verbose_name = "SGSI 4.2 Partes interesadas"
''',
    'apps/context42/models.py': '''from django.conf import settings
from django.db import models
from django.utils.text import slugify


class Context42DocumentKind(models.TextChoices):
    INTERESTED = "interested", "Partes interesadas"
    LEGAL = "legal", "Lista de requisitos legales"


class Context42Document(models.Model):
    slug = models.SlugField(max_length=80, unique=True)
    title = models.CharField(max_length=200)
    kind = models.CharField(max_length=20, choices=Context42DocumentKind.choices)
    description = models.TextField(blank=True)
    sort_order = models.PositiveIntegerField(default=10)
    is_active = models.BooleanField(default=True, db_index=True)
    created_at = models.DateTimeField(auto_now_add=True)
    updated_at = models.DateTimeField(auto_now=True)

    class Meta:
        ordering = ("sort_order", "title")

    def __str__(self):
        return self.title


class Context42DocumentVersion(models.Model):
    document = models.ForeignKey(Context42Document, related_name="versions", on_delete=models.CASCADE)
    file = models.FileField(upload_to="context42/")
    original_name = models.CharField(max_length=255)
    version_label = models.CharField(max_length=40, blank=True)
    notes = models.TextField(blank=True)
    is_current = models.BooleanField(default=True, db_index=True)
    created_at = models.DateTimeField(auto_now_add=True)
    updated_at = models.DateTimeField(auto_now=True)
    created_by = models.ForeignKey(settings.AUTH_USER_MODEL, null=True, blank=True, related_name="context42_versions_created", on_delete=models.SET_NULL)
    updated_by = models.ForeignKey(settings.AUTH_USER_MODEL, null=True, blank=True, related_name="context42_versions_updated", on_delete=models.SET_NULL)

    class Meta:
        ordering = ("-created_at",)

    def __str__(self):
        return f"{self.document.title} - {self.version_label or self.original_name}"

    @property
    def extension(self):
        name = self.original_name or self.file.name
        return name.lower().rsplit('.', 1)[-1] if '.' in name else ''


class InterestedPartyRow(models.Model):
    version = models.ForeignKey(Context42DocumentVersion, related_name="party_rows", on_delete=models.CASCADE)
    source_row = models.PositiveIntegerField(null=True, blank=True)
    party_name = models.CharField(max_length=180, db_index=True)
    need = models.TextField()
    expectation = models.TextField()
    is_active = models.BooleanField(default=True, db_index=True)
    created_at = models.DateTimeField(auto_now_add=True)
    updated_at = models.DateTimeField(auto_now=True)

    class Meta:
        ordering = ("source_row", "party_name", "id")


class LegalRequirementRow(models.Model):
    version = models.ForeignKey(Context42DocumentVersion, related_name="legal_rows", on_delete=models.CASCADE)
    source_row = models.PositiveIntegerField(null=True, blank=True)
    item_no = models.CharField(max_length=30, blank=True)
    requirement = models.TextField()
    promulgated_by = models.CharField(max_length=200, blank=True)
    location = models.TextField(blank=True)
    responsible = models.CharField(max_length=200, blank=True)
    stakeholders = models.TextField(blank=True)
    status = models.CharField(max_length=80, blank=True)
    is_active = models.BooleanField(default=True, db_index=True)
    created_at = models.DateTimeField(auto_now_add=True)
    updated_at = models.DateTimeField(auto_now=True)

    class Meta:
        ordering = ("source_row", "id")
''',
    'apps/context42/forms.py': '''from django import forms

from .models import Context42DocumentVersion


class Context42VersionUploadForm(forms.ModelForm):
    class Meta:
        model = Context42DocumentVersion
        fields = ("file", "version_label", "notes")
        widgets = {
            "notes": forms.Textarea(attrs={"rows": 3}),
        }

    def __init__(self, *args, document=None, **kwargs):
        super().__init__(*args, **kwargs)
        self.document = document
        for field in self.fields.values():
            field.widget.attrs.setdefault("class", "form-control")
        self.fields["file"].help_text = "Sube una nueva versión del archivo. Se conservará el historial y esta versión quedará vigente."''',
    'apps/context42/parser.py': '''from pathlib import Path

from openpyxl import load_workbook


def _clean(value):
    if value is None:
        return ""
    return " ".join(str(value).replace("\xa0", " ").split()).strip()


def parse_interested_parties(file_path):
    wb = load_workbook(filename=file_path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = []
    current_party = ""
    for row_idx in range(4, ws.max_row + 1):
        party = _clean(ws.cell(row=row_idx, column=1).value)
        need = _clean(ws.cell(row=row_idx, column=2).value)
        expectation = _clean(ws.cell(row=row_idx, column=3).value)
        if party:
            current_party = party
        if current_party and (need or expectation):
            rows.append({
                "source_row": row_idx,
                "party_name": current_party,
                "need": need,
                "expectation": expectation,
            })
    return rows


def parse_legal_requirements(file_path):
    wb = load_workbook(filename=file_path, data_only=True)
    ws = wb["Documentos"] if "Documentos" in wb.sheetnames else wb[wb.sheetnames[0]]
    rows = []
    for row_idx in range(4, ws.max_row + 1):
        item_no = _clean(ws.cell(row=row_idx, column=1).value)
        requirement = _clean(ws.cell(row=row_idx, column=2).value)
        promulgated_by = _clean(ws.cell(row=row_idx, column=3).value)
        location = _clean(ws.cell(row=row_idx, column=4).value)
        responsible = _clean(ws.cell(row=row_idx, column=5).value)
        stakeholders = _clean(ws.cell(row=row_idx, column=6).value)
        status = _clean(ws.cell(row=row_idx, column=7).value)
        if not any([item_no, requirement, promulgated_by, location, responsible, stakeholders, status]):
            continue
        if not requirement and not item_no:
            continue
        rows.append({
            "source_row": row_idx,
            "item_no": item_no,
            "requirement": requirement,
            "promulgated_by": promulgated_by,
            "location": location,
            "responsible": responsible,
            "stakeholders": stakeholders,
            "status": status,
        })
    return rows
''',
    'apps/context42/services.py': '''from collections import OrderedDict
from pathlib import Path

from django.db import transaction

from .models import (
    Context42Document,
    Context42DocumentKind,
    Context42DocumentVersion,
    InterestedPartyRow,
    LegalRequirementRow,
)
from .parser import parse_interested_parties, parse_legal_requirements


DEFAULT_DOCUMENTS = [
    {
        "slug": "partes-interesadas",
        "title": "Partes interesadas",
        "kind": Context42DocumentKind.INTERESTED,
        "sort_order": 10,
        "description": "Matriz de partes interesadas con necesidades y expectativas.",
    },
    {
        "slug": "requisitos-legales",
        "title": "Lista de requisitos legales, normativos y contractuales",
        "kind": Context42DocumentKind.LEGAL,
        "sort_order": 20,
        "description": "Listado vigente de requisitos legales, normativos y contractuales aplicables al SGSI.",
    },
]


def ensure_documents():
    docs = []
    for payload in DEFAULT_DOCUMENTS:
        doc, _ = Context42Document.objects.get_or_create(
            slug=payload["slug"],
            defaults=payload,
        )
        changed = False
        for key, value in payload.items():
            if getattr(doc, key) != value:
                setattr(doc, key, value)
                changed = True
        if changed:
            doc.save(update_fields=["title", "kind", "sort_order", "description"])
        docs.append(doc)
    return docs


def detect_version_label(file_name):
    cleaned = (file_name or "").upper()
    for token in cleaned.replace("-", " ").replace("_", " ").split():
        if token.startswith("V") and any(ch.isdigit() for ch in token):
            return token
    return "NUEVA"


@transaction.atomic

def register_version(document, uploaded_file, actor=None, notes="", version_label=""):
    if not version_label:
        version_label = detect_version_label(getattr(uploaded_file, "name", ""))

    Context42DocumentVersion.objects.filter(document=document, is_current=True).update(is_current=False)
    version = Context42DocumentVersion.objects.create(
        document=document,
        file=uploaded_file,
        original_name=getattr(uploaded_file, "name", "archivo"),
        version_label=version_label,
        notes=notes,
        is_current=True,
        created_by=actor if getattr(actor, "is_authenticated", False) else None,
        updated_by=actor if getattr(actor, "is_authenticated", False) else None,
    )
    import_version_rows(version)
    return version


@transaction.atomic

def import_version_rows(version):
    InterestedPartyRow.objects.filter(version=version).delete()
    LegalRequirementRow.objects.filter(version=version).delete()

    path = version.file.path
    if version.document.kind == Context42DocumentKind.INTERESTED:
        for row in parse_interested_parties(path):
            InterestedPartyRow.objects.create(version=version, **row)
    elif version.document.kind == Context42DocumentKind.LEGAL:
        for row in parse_legal_requirements(path):
            LegalRequirementRow.objects.create(version=version, **row)


def party_groups(version):
    groups = OrderedDict()
    if not version:
        return []
    rows = version.party_rows.filter(is_active=True).order_by("source_row", "party_name")
    for row in rows:
        groups.setdefault(row.party_name, []).append(row)
    return [{"name": name, "rows": items} for name, items in groups.items()]


def legal_summary(version):
    if not version:
        return {"statuses": [], "responsibles": [], "promulgated": [], "count": 0}
    rows = list(version.legal_rows.filter(is_active=True))
    return {
        "statuses": sorted({r.status for r in rows if r.status}),
        "responsibles": sorted({r.responsible for r in rows if r.responsible}),
        "promulgated": sorted({r.promulgated_by for r in rows if r.promulgated_by}),
        "count": len(rows),
    }
''',
    'apps/context42/views.py': '''from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied
from django.http import FileResponse, Http404
from django.shortcuts import get_object_or_404, redirect, render
from django.views.decorators.clickjacking import xframe_options_sameorigin

from .forms import Context42VersionUploadForm
from .models import Context42Document, Context42DocumentKind, Context42DocumentVersion
from .services import ensure_documents, legal_summary, party_groups, register_version


def _can_change(request):
    return request.user.is_superuser or request.user.has_perm("context42.add_context42documentversion")


@login_required
def context42_home(request):
    ensure_documents()
    documents = list(Context42Document.objects.filter(is_active=True).prefetch_related("versions").order_by("sort_order"))
    by_slug = {doc.slug: doc for doc in documents}
    for doc in documents:
        doc.current_version = next((v for v in doc.versions.all() if v.is_current), None)
        doc.history = list(doc.versions.all())

    parties_doc = by_slug.get("partes-interesadas")
    legal_doc = by_slug.get("requisitos-legales")
    party_version = getattr(parties_doc, "current_version", None)
    legal_version = getattr(legal_doc, "current_version", None)

    context = {
        "documents": documents,
        "parties_doc": parties_doc,
        "legal_doc": legal_doc,
        "party_groups": party_groups(party_version),
        "party_rows_count": party_version.party_rows.filter(is_active=True).count() if party_version else 0,
        "party_count": len(party_groups(party_version)),
        "legal_rows": list(legal_version.legal_rows.filter(is_active=True).order_by("source_row")) if legal_version else [],
        "legal_summary": legal_summary(legal_version),
        "can_change": _can_change(request),
    }
    return render(request, "context42/context42.html", context)


@login_required
def upload_version(request, slug):
    ensure_documents()
    document = get_object_or_404(Context42Document, slug=slug, is_active=True)
    if not _can_change(request):
        raise PermissionDenied

    if request.method == "POST":
        form = Context42VersionUploadForm(request.POST, request.FILES, document=document)
        if form.is_valid():
            register_version(
                document=document,
                uploaded_file=form.cleaned_data["file"],
                actor=request.user,
                notes=form.cleaned_data.get("notes", ""),
                version_label=form.cleaned_data.get("version_label", ""),
            )
            messages.success(
                request,
                "Nueva versión registrada. La pantalla ya usa el archivo vigente y el historial anterior se conserva.",
            )
            return redirect(f"/sgsi/4.2/?open={document.slug}&updated=1")
    else:
        form = Context42VersionUploadForm(document=document)

    return render(request, "context42/upload_version.html", {"document": document, "form": form})


@login_required
@xframe_options_sameorigin
def download_current(request, slug):
    ensure_documents()
    document = get_object_or_404(Context42Document, slug=slug, is_active=True)
    version = document.versions.filter(is_current=True).first()
    if not version:
        raise Http404("No hay archivo vigente.")
    response = FileResponse(version.file.open("rb"), as_attachment=True, filename=version.original_name)
    return response
''',
    'apps/context42/urls.py': '''from django.urls import path

from . import views

app_name = "context42"

urlpatterns = [
    path("sgsi/4.2/", views.context42_home, name="home"),
    path("sgsi/4.2/<slug:slug>/subir/", views.upload_version, name="upload_version"),
    path("sgsi/4.2/<slug:slug>/descargar/", views.download_current, name="download_current"),
]
''',
    'apps/context42/admin.py': '''from django.contrib import admin

from .models import Context42Document, Context42DocumentVersion, InterestedPartyRow, LegalRequirementRow


@admin.register(Context42Document)
class Context42DocumentAdmin(admin.ModelAdmin):
    list_display = ("title", "kind", "sort_order", "is_active", "updated_at")
    search_fields = ("title", "slug")


@admin.register(Context42DocumentVersion)
class Context42DocumentVersionAdmin(admin.ModelAdmin):
    list_display = ("document", "version_label", "original_name", "is_current", "created_at")
    list_filter = ("document", "is_current")
    search_fields = ("original_name", "version_label", "document__title")


@admin.register(InterestedPartyRow)
class InterestedPartyRowAdmin(admin.ModelAdmin):
    list_display = ("party_name", "source_row", "version")
    search_fields = ("party_name", "need", "expectation")


@admin.register(LegalRequirementRow)
class LegalRequirementRowAdmin(admin.ModelAdmin):
    list_display = ("item_no", "promulgated_by", "responsible", "status")
    search_fields = ("requirement", "promulgated_by", "responsible", "status")
''',
    'apps/context42/migrations/__init__.py': '',
    'templates/context42/context42.html': '''{% extends "base.html" %}
{% load static %}

{% block extra_css %}
<link rel="stylesheet" href="{% static 'css/context42_v70.css' %}">
{% endblock %}

{% block content %}
<div class="context42-page">
    <div class="context42-header">
        <div>
            <div class="breadcrumb small text-muted">SGSI / Cláusula 4.2</div>
            <h1>4.2. Comprender las necesidades y expectativas de las partes interesadas</h1>
            <p>
                La determinación de las partes interesadas que afectan o pueden ser afectadas por el SGSI,
                además de los requisitos legales, se gestiona aquí con versiones vigentes y actualización inmediata.
            </p>
        </div>
        <div class="context42-badge">{{ documents|length }} documentos</div>
    </div>

    <div class="context42-note">
        Cuando subes una nueva versión desde esta misma pantalla, el sistema la toma como vigente inmediatamente.
        La versión anterior no se elimina y queda disponible en el historial.
    </div>

    <div class="context42-accordion">
        <details class="context42-card" {% if request.GET.open == 'partes-interesadas' or not request.GET.open %}open{% endif %} data-document-slug="partes-interesadas">
            <summary>
                <div>
                    <div class="context42-code">01</div>
                    <div>
                        <h2>Partes interesadas</h2>
                        <p>Matriz de necesidades y expectativas.</p>
                    </div>
                </div>
                <div class="context42-summary-right">{{ party_count }} partes / {{ party_rows_count }} filas</div>
            </summary>

            <div class="context42-toolbar">
                <div>
                    {% if parties_doc.current_version %}
                    <strong>Archivo vigente:</strong> {{ parties_doc.current_version.original_name }}
                    <span class="muted">· versión {{ parties_doc.current_version.version_label|default:"SIN-VERSIÓN" }}</span>
                    {% else %}
                    <span class="muted">Todavía no hay archivo cargado.</span>
                    {% endif %}
                </div>
                <div class="context42-actions">
                    {% if parties_doc.current_version %}
                    <a class="btn btn-light" href="{% url 'context42:download_current' parties_doc.slug %}">Descargar</a>
                    {% endif %}
                    {% if can_change %}
                    <a class="btn btn-primary" href="{% url 'context42:upload_version' parties_doc.slug %}">Subir nueva versión</a>
                    {% endif %}
                </div>
            </div>

            {% if party_groups %}
            <div class="context42-stats-grid compact">
                <div class="stat-box"><span>{{ party_count }}</span><small>partes interesadas</small></div>
                <div class="stat-box"><span>{{ party_rows_count }}</span><small>registros importados</small></div>
            </div>
            <div class="parties-groups">
                {% for group in party_groups %}
                <section class="party-group">
                    <div class="party-title">{{ group.name }}</div>
                    <div class="table-wrap">
                        <table class="clean-table">
                            <thead>
                                <tr>
                                    <th>Necesidades</th>
                                    <th>Expectativas</th>
                                </tr>
                            </thead>
                            <tbody>
                                {% for row in group.rows %}
                                <tr>
                                    <td>{{ row.need }}</td>
                                    <td>{{ row.expectation }}</td>
                                </tr>
                                {% endfor %}
                            </tbody>
                        </table>
                    </div>
                </section>
                {% endfor %}
            </div>
            {% else %}
            <div class="empty-panel">No hay información cargada todavía. Sube el archivo Excel de partes interesadas para ver la matriz aquí mismo.</div>
            {% endif %}

            {% if parties_doc.history %}
            <div class="history-box">
                <h3>Historial</h3>
                <ul>
                    {% for item in parties_doc.history %}
                    <li>
                        {{ item.original_name }}
                        <span class="muted">· {{ item.version_label|default:"SIN-VERSIÓN" }} · {{ item.created_at|date:"d/m/Y H:i" }}{% if item.is_current %} · vigente{% endif %}</span>
                    </li>
                    {% endfor %}
                </ul>
            </div>
            {% endif %}
        </details>

        <details class="context42-card" {% if request.GET.open == 'requisitos-legales' %}open{% endif %} data-document-slug="requisitos-legales">
            <summary>
                <div>
                    <div class="context42-code">02</div>
                    <div>
                        <h2>Lista de requisitos legales, normativos y contractuales</h2>
                        <p>Listado vigente aplicable al SGSI.</p>
                    </div>
                </div>
                <div class="context42-summary-right">{{ legal_summary.count }} requisitos</div>
            </summary>

            <div class="context42-toolbar">
                <div>
                    {% if legal_doc.current_version %}
                    <strong>Archivo vigente:</strong> {{ legal_doc.current_version.original_name }}
                    <span class="muted">· versión {{ legal_doc.current_version.version_label|default:"SIN-VERSIÓN" }}</span>
                    {% else %}
                    <span class="muted">Todavía no hay archivo cargado.</span>
                    {% endif %}
                </div>
                <div class="context42-actions">
                    {% if legal_doc.current_version %}
                    <a class="btn btn-light" href="{% url 'context42:download_current' legal_doc.slug %}">Descargar</a>
                    {% endif %}
                    {% if can_change %}
                    <a class="btn btn-primary" href="{% url 'context42:upload_version' legal_doc.slug %}">Subir nueva versión</a>
                    {% endif %}
                </div>
            </div>

            {% if legal_rows %}
            <div class="context42-stats-grid">
                <div class="stat-box"><span>{{ legal_summary.count }}</span><small>requisitos</small></div>
                <div class="stat-box"><span>{{ legal_summary.statuses|length }}</span><small>estados</small></div>
                <div class="stat-box"><span>{{ legal_summary.responsibles|length }}</span><small>responsables</small></div>
                <div class="stat-box"><span>{{ legal_summary.promulgated|length }}</span><small>entes emisores</small></div>
            </div>
            <div class="table-wrap table-wrap-wide">
                <table class="clean-table clean-table-legal">
                    <thead>
                        <tr>
                            <th>N°</th>
                            <th>Requisito</th>
                            <th>Promulgada por</th>
                            <th>Ubicación</th>
                            <th>Responsable</th>
                            <th>Partes interesadas</th>
                            <th>Estado</th>
                        </tr>
                    </thead>
                    <tbody>
                        {% for row in legal_rows %}
                        <tr>
                            <td>{{ row.item_no }}</td>
                            <td>{{ row.requirement }}</td>
                            <td>{{ row.promulgated_by }}</td>
                            <td>{{ row.location }}</td>
                            <td>{{ row.responsible }}</td>
                            <td>{{ row.stakeholders }}</td>
                            <td><span class="status-pill">{{ row.status }}</span></td>
                        </tr>
                        {% endfor %}
                    </tbody>
                </table>
            </div>
            {% else %}
            <div class="empty-panel">No hay información cargada todavía. Sube el archivo Excel de requisitos legales para ver la tabla aquí mismo.</div>
            {% endif %}

            {% if legal_doc.history %}
            <div class="history-box">
                <h3>Historial</h3>
                <ul>
                    {% for item in legal_doc.history %}
                    <li>
                        {{ item.original_name }}
                        <span class="muted">· {{ item.version_label|default:"SIN-VERSIÓN" }} · {{ item.created_at|date:"d/m/Y H:i" }}{% if item.is_current %} · vigente{% endif %}</span>
                    </li>
                    {% endfor %}
                </ul>
            </div>
            {% endif %}
        </details>
    </div>
</div>
{% endblock %}
''',
    'templates/context42/upload_version.html': '''{% extends "base.html" %}

{% block content %}
<div class="container" style="max-width: 760px; padding-top: 32px;">
    <div class="card">
        <div class="card-body">
            <h1 style="font-size: 28px; margin-bottom: 8px;">Actualizar archivo</h1>
            <p style="color: #6b7280; margin-bottom: 20px;">{{ document.title }}</p>
            <form method="post" enctype="multipart/form-data">
                {% csrf_token %}
                {% for field in form %}
                <div class="form-group" style="margin-bottom: 16px;">
                    <label style="display:block; font-weight:600; margin-bottom:6px;">{{ field.label }}</label>
                    {{ field }}
                    {% if field.help_text %}<small class="form-text text-muted">{{ field.help_text }}</small>{% endif %}
                    {% if field.errors %}<div class="text-danger">{{ field.errors }}</div>{% endif %}
                </div>
                {% endfor %}
                <div style="display:flex; gap:10px; margin-top: 18px;">
                    <button type="submit" class="btn btn-primary">Guardar nueva versión</button>
                    <a href="/sgsi/4.2/?open={{ document.slug }}" class="btn btn-light">Cancelar</a>
                </div>
            </form>
        </div>
    </div>
</div>
{% endblock %}
''',
    'static/css/context42_v70.css': '''.context42-page{padding:28px 22px 32px}.context42-header{display:flex;justify-content:space-between;gap:18px;align-items:flex-start;margin-bottom:18px}.context42-header h1{font-size:40px;line-height:1.15;color:#243b53;margin:2px 0 10px}.context42-header p{font-size:16px;color:#52606d;max-width:980px;margin:0}.context42-badge{background:#eef4fb;border:1px solid #d3e2f2;color:#355070;padding:8px 14px;border-radius:999px;font-size:13px;white-space:nowrap}.context42-note{background:#f8fbfd;border-left:4px solid #8aaec8;color:#4a6177;padding:12px 14px;border-radius:6px;margin-bottom:16px}.context42-accordion{display:grid;gap:16px}.context42-card{background:#fff;border:1px solid #dbe7f1;border-radius:14px;overflow:hidden}.context42-card summary{list-style:none;display:flex;align-items:center;justify-content:space-between;gap:16px;padding:18px 20px;cursor:pointer}.context42-card summary::-webkit-details-marker{display:none}.context42-card summary>div:first-child{display:flex;align-items:flex-start;gap:14px}.context42-code{width:40px;height:40px;border-radius:10px;background:#edf3f8;color:#355070;display:flex;align-items:center;justify-content:center;font-weight:700}.context42-card h2{font-size:23px;color:#243b53;margin:0 0 4px}.context42-card summary p{margin:0;color:#6b7c8f;font-size:14px}.context42-summary-right{font-size:13px;color:#6b7c8f;white-space:nowrap}.context42-toolbar{display:flex;justify-content:space-between;gap:16px;align-items:center;border-top:1px solid #edf2f7;padding:14px 20px 0}.context42-actions{display:flex;gap:10px;flex-wrap:wrap}.muted{color:#7b8794}.context42-stats-grid{display:grid;grid-template-columns:repeat(4,minmax(0,1fr));gap:12px;padding:16px 20px 0}.context42-stats-grid.compact{grid-template-columns:repeat(2,minmax(0,1fr))}.stat-box{background:#f8fafc;border:1px solid #e5edf5;border-radius:12px;padding:14px 16px}.stat-box span{display:block;font-size:26px;font-weight:700;color:#243b53;line-height:1}.stat-box small{display:block;margin-top:6px;color:#6b7c8f;font-size:13px}.parties-groups{padding:16px 20px 18px;display:grid;gap:16px}.party-group{border:1px solid #e6eef5;border-radius:12px;overflow:hidden}.party-title{padding:12px 14px;font-weight:700;color:#243b53;background:#f8fbfd;border-bottom:1px solid #e6eef5}.table-wrap{overflow:auto}.table-wrap-wide{padding:16px 20px 18px}.clean-table{width:100%;border-collapse:collapse}.clean-table th,.clean-table td{padding:12px 14px;border-bottom:1px solid #ecf1f5;vertical-align:top;text-align:left}.clean-table th{font-size:13px;text-transform:none;color:#5b7083;background:#fbfdff}.clean-table td{font-size:14px;color:#334e68;line-height:1.45}.clean-table-legal th:nth-child(1),.clean-table-legal td:nth-child(1){white-space:nowrap;width:60px}.status-pill{display:inline-block;padding:4px 10px;border-radius:999px;background:#edf7ed;color:#2d6a4f;font-size:12px;border:1px solid #cfe8cf}.history-box{padding:0 20px 20px}.history-box h3{margin:6px 0 10px;font-size:15px;color:#243b53}.history-box ul{margin:0;padding-left:18px;color:#52606d}.history-box li{margin-bottom:5px}.empty-panel{padding:18px 20px 20px;color:#6b7c8f}.breadcrumb{margin-bottom:6px;color:#8a9aa8}@media (max-width: 980px){.context42-header,.context42-toolbar{flex-direction:column;align-items:flex-start}.context42-header h1{font-size:30px}.context42-stats-grid,.context42-stats-grid.compact{grid-template-columns:1fr 1fr}}@media (max-width: 640px){.context42-page{padding:22px 14px 28px}.context42-card summary{padding:16px}.context42-stats-grid,.context42-stats-grid.compact{grid-template-columns:1fr}}
''',
}


def write_file(rel_path, content):
    path = ROOT / rel_path
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        backup = path.with_name(path.name + '.before_context42_v70.bak')
        if not backup.exists():
            backup.write_text(path.read_text(encoding='utf-8'), encoding='utf-8')
    path.write_text(textwrap.dedent(content).lstrip('\n'), encoding='utf-8')
    print('OK:', rel_path)


def patch_settings():
    candidates = [ROOT / 'config/settings.py', ROOT / 'config/settings/base.py']
    for path in candidates:
        if path.exists():
            text = path.read_text(encoding='utf-8')
            if 'apps.context42.apps.Context42Config' not in text:
                if 'INSTALLED_APPS = [' in text:
                    text = text.replace('INSTALLED_APPS = [', 'INSTALLED_APPS = [\n    "apps.context42.apps.Context42Config",', 1)
                elif 'INSTALLED_APPS += [' in text:
                    text = text.replace('INSTALLED_APPS += [', 'INSTALLED_APPS += [\n    "apps.context42.apps.Context42Config",', 1)
                else:
                    continue
                path.write_text(text, encoding='utf-8')
                print('OK:', path.as_posix(), '(Context42Config agregado)')
            return
    print('ADVERTENCIA: no se pudo ubicar settings.py para agregar apps.context42')


def patch_urls():
    path = ROOT / 'config/urls.py'
    if not path.exists():
        print('ADVERTENCIA: no se encontró config/urls.py')
        return
    text = path.read_text(encoding='utf-8')
    if 'apps.context42.urls' in text:
        print('OK: config/urls.py ya incluye apps.context42.urls')
        return
    marker = 'from django.urls import include, path'
    if marker not in text:
        print('ADVERTENCIA: no se encontró import base en config/urls.py')
        return
    if 'path("", include("apps.dashboard.urls"))' in text:
        replacement = 'path("", include("apps.context42.urls")),\n    path("", include("apps.dashboard.urls"))'
        text = text.replace('path("", include("apps.dashboard.urls"))', replacement, 1)
    elif 'path("", include("apps.context41.urls"))' in text:
        replacement = 'path("", include("apps.context42.urls")),\n    path("", include("apps.context41.urls"))'
        text = text.replace('path("", include("apps.context41.urls"))', replacement, 1)
    elif 'urlpatterns = [' in text:
        text = text.replace('urlpatterns = [', 'urlpatterns = [\n    path("", include("apps.context42.urls")),', 1)
    else:
        print('ADVERTENCIA: no se pudo insertar la ruta de context42')
        return
    path.write_text(text, encoding='utf-8')
    print('OK: config/urls.py (ruta /sgsi/4.2/ agregada)')


def patch_base_menu():
    path = ROOT / 'templates/base.html'
    if not path.exists():
        return
    text = path.read_text(encoding='utf-8')
    changed = False
    if '/sgsi/4.2/' not in text and '4.2' in text:
        text = text.replace('href="/sgsi/4.1/"', 'href="/sgsi/4.1/"', 1)
    clause_link_old = 'href="/sgsi/4.2"'
    if clause_link_old in text:
        text = text.replace(clause_link_old, 'href="/sgsi/4.2/"')
        changed = True
    if changed:
        path.write_text(text, encoding='utf-8')
        print('OK: templates/base.html')


def main():
    print('=== INSTALACIÓN CONTEXTO 4.2 V70 ===')
    for rel_path, content in FILES.items():
        write_file(rel_path, content)
    patch_settings()
    patch_urls()
    patch_base_menu()
    print('')
    print('Contexto 4.2 instalado.')
    print('Características:')
    print('- Vista elegante y simple para las 2 evidencias de la cláusula 4.2.')
    print('- Subir nueva versión desde la misma pantalla.')
    print('- La nueva versión queda vigente inmediatamente.')
    print('- El historial anterior se conserva.')
    print('- El Excel de Partes Interesadas se importa y se muestra agrupado por parte interesada.')
    print('- El Excel de requisitos legales se importa y se muestra en tabla completa.')
    print('')
    print('Siguiente paso:')
    print('1) python -m py_compile apps\\context42\\models.py apps\\context42\\forms.py apps\\context42\\parser.py apps\\context42\\services.py apps\\context42\\views.py apps\\context42\\urls.py')
    print('2) docker compose exec web python manage.py makemigrations context42')
    print('3) docker compose exec web python manage.py migrate')
    print('4) docker compose exec web python manage.py check')
    print('5) docker compose restart web')
    print('6) entra a http://localhost:8000/sgsi/4.2/')
    print('7) sube los archivos:')
    print('   - 07 - Partes Interesadas_V0.4.xlsx')
    print('   - 03 - Lista_de_requisitos_legales_normativos_contractuales_V0.15.xlsm')


if __name__ == '__main__':
    main()
