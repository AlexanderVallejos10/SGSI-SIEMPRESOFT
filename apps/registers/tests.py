import io
import json

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from .models import RegisterEntry


def plan_siempresoft(year=2026):
    """Excel con el formato del plan de SiempreSoft: título, perfil combinado y meses con X."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Capacitación"
    ws["B1"] = f"PLAN DE CAPACITACIÓN Y CONCIENCIACIÓN {year}"
    ws["A3"], ws["B3"], ws["C3"], ws["O3"] = "PERFIL DE PUESTO", "CONOCIMIENTOS Y HABILIDADES", "MES", "Modalidad"
    for i, letter in enumerate("EFMAMJJASOND"):
        ws.cell(row=4, column=3 + i, value=letter)
    ws["A5"], ws["B5"], ws["C5"], ws["O5"] = "Desarrollador", "Open ID – Connect", "X", "Capacitación externa"
    ws["B6"], ws["D6"], ws["O6"] = "Curso de SQL", "X", "Capacitación externa"
    ws.merge_cells("A5:A6")
    ws["A7"], ws["B7"], ws["N7"], ws["O7"] = "Todas las áreas", "Ciberseguridad defensiva", "X", "Capacitación externa"
    ws["B9"] = "Nota: todas las capacitaciones serán evaluadas"
    out = io.BytesIO()
    wb.save(out)
    return SimpleUploadedFile(f"Plan de capacitación y concienciación {year}.xlsx", out.getvalue())


class RegistrosDelSGSI(TestCase):
    @classmethod
    def setUpTestData(cls):
        User = get_user_model()
        cls.admin = User.objects.create_user(username="reg-admin", business_code="T-REG", is_superuser=True)
        cls.reader = User.objects.create_user(username="reg-reader", business_code="T-REG-R")

    def setUp(self):
        self.client.force_login(self.admin)

    def test_importa_el_plan_con_el_formato_de_siempresoft(self):
        url = reverse("registers:import", args=["plan-capacitacion"])
        response = self.client.post(url, {"file": plan_siempresoft(), "mode": "replace"})
        self.assertRedirects(response, reverse("registers:detail", args=["plan-capacitacion"]) + "?anio=2026", fetch_redirect_response=False)
        rows = list(RegisterEntry.objects.filter(register="plan-capacitacion", year=2026).order_by("order"))
        self.assertEqual(len(rows), 3)  # la nota al pie no entra
        self.assertEqual(rows[1].data["perfil"], "Desarrollador")  # celda combinada
        self.assertEqual(rows[1].data["meses"], [2])
        self.assertEqual(rows[2].data["meses"], [12])
        page = self.client.get(reverse("registers:detail", args=["plan-capacitacion"]) + "?anio=2026")
        self.assertContains(page, "Open ID – Connect")
        self.assertContains(page, "Temas programados")

    def test_editar_meses_y_estado_actualiza_indicadores(self):
        self.client.post(reverse("registers:import", args=["plan-capacitacion"]), {"file": plan_siempresoft()})
        entry = RegisterEntry.objects.filter(register="plan-capacitacion").order_by("order").first()
        response = self.client.post(
            reverse("registers:update_row", args=[entry.pk]),
            data=json.dumps({"meses": [1, 3], "estado": "Realizada"}),
            content_type="application/json",
        )
        data = response.json()
        entry.refresh_from_db()
        self.assertEqual(entry.data["meses"], [1, 3])
        self.assertEqual(data["summary"]["cards"][1]["value"], 1)  # realizados

    def test_descarga_en_excel_y_vuelve_a_importar_igual(self):
        self.client.post(reverse("registers:import", args=["plan-capacitacion"]), {"file": plan_siempresoft()})
        before = [e.data for e in RegisterEntry.objects.filter(register="plan-capacitacion").order_by("order")]
        response = self.client.get(reverse("registers:export", args=["plan-capacitacion"]) + "?anio=2026")
        self.assertEqual(response["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        wb = load_workbook(io.BytesIO(response.content))
        self.assertIn("Capacitación", wb.sheetnames)
        again = SimpleUploadedFile("descargado.xlsx", response.content)
        self.client.post(reverse("registers:import", args=["plan-capacitacion"]), {"file": again, "year": "2026"})
        after = [e.data for e in RegisterEntry.objects.filter(register="plan-capacitacion").order_by("order")]
        self.assertEqual(before, after)

    def test_agregar_y_quitar_filas(self):
        response = self.client.post(
            reverse("registers:create_row", args=["software-autorizado"]),
            data=json.dumps({"section": "software", "data": {"software": "7-Zip", "estado": "en uso"}}),
            content_type="application/json",
        )
        self.assertIn("7-Zip", response.json()["html"])
        entry = RegisterEntry.objects.get(register="software-autorizado")
        self.assertEqual(entry.data["estado"], "En uso")
        self.client.post(reverse("registers:delete_row", args=[entry.pk]))
        self.assertFalse(RegisterEntry.objects.exists())

    def test_quien_no_tiene_permiso_solo_puede_ver(self):
        self.client.force_login(self.reader)
        self.assertEqual(self.client.get(reverse("registers:index")).status_code, 200)
        page = self.client.get(reverse("registers:detail", args=["obligaciones-ose"]))
        self.assertNotContains(page, "data-open-import")
        response = self.client.post(
            reverse("registers:create_row", args=["obligaciones-ose"]),
            data=json.dumps({"section": "requisitos", "data": {"obligacion": "x"}}),
            content_type="application/json",
        )
        self.assertEqual(response.status_code, 403)

    def test_archivo_que_no_corresponde_avisa_sin_romper(self):
        wb = Workbook()
        wb.active["A1"] = "otra cosa"
        out = io.BytesIO()
        wb.save(out)
        response = self.client.post(
            reverse("registers:import", args=["software-autorizado"]),
            {"file": SimpleUploadedFile("cualquiera.xlsx", out.getvalue())}, follow=True,
        )
        self.assertContains(response, "No se reconocieron filas")
        self.assertFalse(RegisterEntry.objects.exists())


class RegistrosDeTablaGenerica(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.admin = get_user_model().objects.create_user(username="reg-tab", business_code="T-REG-T", is_superuser=True)

    def setUp(self):
        self.client.force_login(self.admin)

    def test_librerias_se_leen_por_encabezado_y_se_muestran(self):
        wb = Workbook()
        ws = wb.active
        ws["B1"] = "INVENTARIO DE LIBRERÍAS EXTERNAS AUTORIZADAS"
        ws["A2"], ws["B2"] = "NOMBRE DE LIBRERÍA", "PROVEEDOR/AUTOR"
        ws["A3"], ws["B3"] = "Stimulsoft.Base", "Stimulsoft"
        ws["A4"], ws["B4"] = "Newtonsoft.Json", "James Newton-King"
        out = io.BytesIO()
        wb.save(out)
        self.client.post(
            reverse("registers:import", args=["librerias-externas"]),
            {"file": SimpleUploadedFile("Inventario de librerías externas autorizadas.xlsx", out.getvalue())},
        )
        rows = RegisterEntry.objects.filter(register="librerias-externas").order_by("order")
        self.assertEqual([r.data["libreria"] for r in rows], ["Stimulsoft.Base", "Newtonsoft.Json"])
        self.assertEqual(rows[0].data["estado"], "Autorizada")
        page = self.client.get(reverse("registers:detail", args=["librerias-externas"]))
        self.assertContains(page, "Newtonsoft.Json")
        self.assertContains(page, "Proveedor o autor")

    def test_fecha_escrita_como_texto_se_guarda_ordenable(self):
        response = self.client.post(
            reverse("registers:create_row", args=["creacion-usuarios"]),
            data=json.dumps({"section": "usuarios", "year": 2021, "data": {"colaborador": "Oscar", "fecha": "27/04/2021"}}),
            content_type="application/json",
        )
        self.assertIn("27/04/2021", response.json()["html"])
        self.assertEqual(RegisterEntry.objects.get(register="creacion-usuarios").data["fecha"], "2021-04-27")


def acta_borrado_docx(soporte, fecha, metodo="Destrucción física"):
    from docx import Document

    doc = Document()
    doc.add_paragraph("ACTA DE BORRADO Y DESTRUCCIÓN DE REGISTROS")
    table = doc.add_table(rows=4, cols=2)
    for row, (label, value) in zip(table.rows, [
        ("Datos sobre los soportes", soporte),
        ("Fecha de borrado o destrucción", fecha),
        ("Método de borrado o destrucción", metodo),
        ("Persona que realizó el proceso (Nombre y cargo)", "Ana Karim Salazar – Oficial de seguridad"),
    ]):
        row.cells[0].text, row.cells[1].text = label, value
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


class RegistrosDelAreaDeSeguridad(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.admin = get_user_model().objects.create_user(username="reg-sec", business_code="T-REG-S", is_superuser=True)

    def setUp(self):
        self.client.force_login(self.admin)

    def test_zip_de_actas_se_reparte_por_el_anio_de_cada_acta(self):
        import zipfile

        out = io.BytesIO()
        with zipfile.ZipFile(out, "w") as zf:
            zf.writestr("Actas de Borrado/2021/01 - ACTA_blinares.docx", acta_borrado_docx("Laptop de B. Linares", "15/03/2021"))
            zf.writestr("Actas de Borrado/2026/03 - ACTA_discos.docx", acta_borrado_docx("Unidades de disco HDD", "14/05/2026"))
            zf.writestr("Actas de Borrado/2026/03 - ACTA_discos - Copia.docx", acta_borrado_docx("Duplicado", "14/05/2026"))
        self.client.post(
            reverse("registers:import", args=["actas-borrado"]),
            {"file": SimpleUploadedFile("Actas de Borrado.zip", out.getvalue())},
        )
        by_year = {e.year: e.data for e in RegisterEntry.objects.filter(register="actas-borrado")}
        self.assertEqual(set(by_year), {2021, 2026})  # la copia no entra
        self.assertEqual(by_year[2026]["fecha"], "2026-05-14")
        self.assertEqual(by_year[2026]["metodo"], "Destrucción física")
        page = self.client.get(reverse("registers:detail", args=["actas-borrado"]) + "?anio=2021")
        self.assertContains(page, "Laptop de B. Linares")

    def test_registro_de_claves_nunca_guarda_el_contenido_de_las_celdas(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "2022"
        for i, month in enumerate(["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio",
                                   "Agosto", "Setiembre", "Octubre", "Noviembre", "Diciembre"]):
            ws.cell(row=4, column=2 + i, value=month)
        ws["A5"], ws["C5"], ws["F5"] = "rraymundo - SUPREMO", "Xk9#claveSecreta2022!", "x"
        out = io.BytesIO()
        wb.save(out)
        self.client.post(
            reverse("registers:import", args=["cambio-claves"]),
            {"file": SimpleUploadedFile("Registro de cambio de claves.xlsx", out.getvalue())},
        )
        entry = RegisterEntry.objects.get(register="cambio-claves")
        self.assertEqual(entry.year, 2022)
        self.assertEqual(entry.data["cuenta"], "rraymundo")
        self.assertEqual(entry.data["tipo"], "Supremo")
        self.assertEqual(entry.data["meses"], [2, 5])
        self.assertNotIn("claveSecreta", json.dumps(entry.data))
        exported = load_workbook(io.BytesIO(
            self.client.get(reverse("registers:export", args=["cambio-claves"]) + "?anio=2022").content
        ))
        cells = [str(c.value) for ws in exported.worksheets for row in ws.iter_rows() for c in row if c.value]
        self.assertFalse(any("claveSecreta" in c for c in cells))


class RegistrosDeVariasHojas(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.admin = get_user_model().objects.create_user(username="reg-multi", business_code="T-REG-M", is_superuser=True)

    def setUp(self):
        self.client.force_login(self.admin)

    def _medidas(self):
        wb = Workbook()
        wb.remove(wb.active)
        for year in ("2025", "2026"):  # la hoja 2026 es copia de la 2025, como en el Excel real
            ws = wb.create_sheet(year)
            ws["C1"] = "SISTEMA DE GESTION DE SEGURIDAD"
            for col, label in enumerate(["Nro", "Fecha de identificación", "Tipo", "Breve descripción / no conformidad",
                                         "Persona responsable de la acción", "Estado de la Medida Correctiva"], start=1):
                ws.cell(row=5, column=col, value=label)
            ws.append([1, "15/01/2025", "Reporte de incidente", "El usuario no podía conectarse", "Jefe de Producción", "Implementado"])
            ws.append([2, "03/02/2028", "Auditoría", "Obs. 1 con fecha mal digitada", "Oficial de Seguridad", "En proceso"])
        out = io.BytesIO()
        wb.save(out)
        return SimpleUploadedFile("01 - Registro centralizado de Medidas correctivas y mejoras.xlsx", out.getvalue())

    def test_hoja_copiada_no_duplica_y_fecha_imposible_usa_el_anio_de_la_hoja(self):
        self.client.post(reverse("registers:import", args=["medidas-correctivas"]), {"file": self._medidas()})
        rows = RegisterEntry.objects.filter(register="medidas-correctivas")
        self.assertEqual(rows.count(), 2)
        self.assertEqual(sorted(r.year for r in rows), [2025, 2025])
        page = self.client.get(reverse("registers:detail", args=["medidas-correctivas"]) + "?anio=2025")
        self.assertContains(page, "Implementadas")

    def test_permisos_guardan_el_ambito_de_cada_hoja(self):
        wb = Workbook()
        wb.remove(wb.active)
        for sheet, name in (("Supremo", "Sheyla Quevedo"), ("VPN", "Elio Mondragón")):
            ws = wb.create_sheet(sheet)
            ws["A4"], ws["B4"], ws["C4"] = "COLABORADOR", "ÁREA / PERFIL DE PUESTO", "FECHA DE PERMISO"
            ws["A5"], ws["B5"], ws["C5"] = name, "Soporte", "01/03/2024"
        out = io.BytesIO()
        wb.save(out)
        self.client.post(reverse("registers:import", args=["registro-permisos"]),
                         {"file": SimpleUploadedFile("16 - Registro de permisos.xlsx", out.getvalue())})
        ambitos = dict(RegisterEntry.objects.filter(register="registro-permisos").values_list("data__colaborador", "data__ambito"))
        self.assertEqual(ambitos, {"Sheyla Quevedo": "Supremo", "Elio Mondragón": "VPN"})

    def test_indice_agrupa_los_registros(self):
        page = self.client.get(reverse("registers:index"))
        self.assertContains(page, "Control de acceso")
        self.assertContains(page, "Registro de incidentes de seguridad de la información")
        self.assertContains(page, "Incidentes y mejora")
