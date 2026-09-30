"""Descarga de cada registro en Excel con el formato de SiempreSoft: título, recuadro de
clasificación, encabezados y, en el plan, los meses marcados con X como en el original.
El mismo archivo sirve de plantilla para volver a importar."""

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .schemas import MONTH_LETTERS

THIN = Side(style="thin", color="7F7F7F")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="D9D9D9")
GROUP_FILL = PatternFill("solid", fgColor="F2F2F2")
TITLE_FONT = Font(name="Calibri", size=14, bold=True)
HEAD_FONT = Font(name="Calibri", size=11, bold=True)
BODY_FONT = Font(name="Calibri", size=11)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
CLASSIFICATION = "DOCUMENTO PARA: SGSI\nÁREA: SEGURIDAD DE LA INFORMACIÓN\nNIVEL DE CONFIDENCIALIDAD: USO INTERNO"


def _title_block(ws, title, last_col):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(1, last_col - 2))
    ws.cell(row=1, column=1, value=title).font = TITLE_FONT
    ws.cell(row=1, column=1).alignment = Alignment(vertical="center")
    ws.merge_cells(start_row=1, start_column=last_col - 1, end_row=1, end_column=last_col)
    cell = ws.cell(row=1, column=last_col - 1, value=CLASSIFICATION)
    cell.font = Font(name="Calibri", size=8)
    cell.alignment = CENTER
    ws.row_dimensions[1].height = 48


def _plan_sheet(ws, title, section, entries):
    capacitacion = section["key"] == "capacitacion"
    first = ["PERFIL DE PUESTO", "CONOCIMIENTOS Y HABILIDADES"] if capacitacion else ["TEMA"]
    month_col = len(first) + 1
    last_col = month_col + 12 + 1  # meses + modalidad + estado
    _title_block(ws, title, last_col)
    head = 3
    for i, label in enumerate(first, start=1):
        ws.merge_cells(start_row=head, start_column=i, end_row=head + 1, end_column=i)
        ws.cell(row=head, column=i, value=label)
    ws.merge_cells(start_row=head, start_column=month_col, end_row=head, end_column=month_col + 11)
    ws.cell(row=head, column=month_col, value="MES")
    for i, letter in enumerate(MONTH_LETTERS):
        ws.cell(row=head + 1, column=month_col + i, value=letter)
    for offset, label in ((12, "Modalidad"), (13, "Estado")):
        ws.merge_cells(start_row=head, start_column=month_col + offset, end_row=head + 1, end_column=month_col + offset)
        ws.cell(row=head, column=month_col + offset, value=label)
    for r in (head, head + 1):
        for c in range(1, last_col + 1):
            cell = ws.cell(row=r, column=c)
            cell.font, cell.fill, cell.border, cell.alignment = HEAD_FONT, HEAD_FILL, BORDER, CENTER
    row = head + 2
    for e in entries:
        d = e.data
        values = ([d.get("perfil", ""), d.get("tema", "")] if capacitacion else [d.get("tema", "")])
        for i, v in enumerate(values, start=1):
            ws.cell(row=row, column=i, value=v)
        for m in d.get("meses", []):
            ws.cell(row=row, column=month_col + m - 1, value="X")
        ws.cell(row=row, column=month_col + 12, value=d.get("modalidad", ""))
        ws.cell(row=row, column=month_col + 13, value=d.get("estado", ""))
        for c in range(1, last_col + 1):
            cell = ws.cell(row=row, column=c)
            cell.font, cell.border = BODY_FONT, BORDER
            cell.alignment = CENTER if month_col <= c < month_col + 12 else WRAP
        row += 1
    widths = ([28, 52] if capacitacion else [60]) + [4] * 12 + [22, 14]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=head + 2, column=month_col)


def _table_sheet(ws, title, section, entries):
    fields = [f for f in section["fields"]]
    last_col = len(fields)
    _title_block(ws, title, last_col)
    head = 3
    for i, f in enumerate(fields, start=1):
        cell = ws.cell(row=head, column=i, value=f["label"].upper())
        cell.font, cell.fill, cell.border, cell.alignment = HEAD_FONT, HEAD_FILL, BORDER, CENTER
    row = head + 1
    for e in entries:
        grupo = e.data.get("tipo") == "grupo"
        for i, f in enumerate(fields, start=1):
            value = e.data.get(f["key"], "")
            if f["type"] == "months":
                value = ", ".join(MONTH_LETTERS[m - 1] for m in value)
            cell = ws.cell(row=row, column=i, value=value)
            cell.font = Font(name="Calibri", size=11, bold=grupo)
            cell.border, cell.alignment = BORDER, WRAP
            if grupo:
                cell.fill = GROUP_FILL
        row += 1
    for i, f in enumerate(fields, start=1):
        ws.column_dimensions[get_column_letter(i)].width = 60 if f["type"] == "longtext" else 22
    ws.freeze_panes = ws.cell(row=head + 1, column=1)


def build(slug, schema, entries_by_section, year=None):
    wb = Workbook()
    wb.remove(wb.active)
    title = schema["title"].format(year=year or "").strip()
    for section in schema["sections"]:
        ws = wb.create_sheet(section.get("sheet") or section["label"][:31])
        entries = entries_by_section.get(section["key"], [])
        if schema["layout"] == "timeline":
            _plan_sheet(ws, title, section, entries)
        else:
            _table_sheet(ws, title, section, entries)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
