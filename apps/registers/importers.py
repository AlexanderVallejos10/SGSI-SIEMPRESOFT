"""Lectura de los Excel de SiempreSoft tal como la empresa los usa: con títulos, celdas combinadas
y meses marcados con X. Cada importador devuelve filas listas para guardar: (sección, datos)."""

import datetime
import io
import re
import unicodedata
import zipfile

from openpyxl import load_workbook

from .schemas import clean_value, get


def _norm(value):
    value = unicodedata.normalize("NFKD", str(value or ""))
    value = "".join(c for c in value if not unicodedata.combining(c)).casefold()
    return " ".join(re.sub(r"[^a-z0-9%?¿]+", " ", value).split())


def _text(value):
    if value is None:
        return ""
    if isinstance(value, (datetime.date, datetime.datetime)):
        return value.strftime("%Y-%m-%d")
    return " ".join(str(value).replace("_x000D_", " ").split())


def _marked(value):
    return str(value or "").strip().lower() in {"x", "✓", "si", "sí", "1"}


def year_from_name(name):
    found = re.findall(r"(20\d\d)", name or "")
    return int(found[-1]) if found else None


def _find_row(ws, predicate, limit=15):
    for row in ws.iter_rows(min_row=1, max_row=limit):
        cells = [c for c in row if c.value not in (None, "")]
        if cells and predicate([_norm(c.value) for c in cells]):
            return row[0].row, {_norm(c.value): c.column for c in cells}
    return None, {}


def _merged_value(ws, row, col):
    """Valor de una celda aunque sea parte de un rango combinado (perfil de puesto)."""
    cell = ws.cell(row=row, column=col)
    if cell.value not in (None, ""):
        return cell.value
    for rng in ws.merged_cells.ranges:
        if rng.min_row <= row <= rng.max_row and rng.min_col <= col <= rng.max_col:
            return ws.cell(row=rng.min_row, column=rng.min_col).value
    return None


# ---------------------------------------------------------------- plan de capacitación

def _plan_sheet(ws, section_key):
    header_row, cols = _find_row(ws, lambda vals: "mes" in vals and ("tema" in vals or "perfil de puesto" in vals))
    if not header_row:
        return []
    month_col = cols["mes"]
    modal_col = next((c for k, c in cols.items() if k.startswith("modalidad")), None)
    text_key = "conocimientos y habilidades" if section_key == "capacitacion" else "tema"
    text_col = next((c for k, c in cols.items() if k.startswith(text_key)), None) or (2 if section_key == "capacitacion" else 1)
    perfil_col = cols.get("perfil de puesto")
    rows = []
    last_perfil = ""
    for r in range(header_row + 2, ws.max_row + 1):
        tema = _text(ws.cell(row=r, column=text_col).value)
        if not tema:
            continue
        months = [i + 1 for i in range(12) if _marked(ws.cell(row=r, column=month_col + i).value)]
        modalidad = _text(ws.cell(row=r, column=modal_col).value) if modal_col else ""
        if not months and not modalidad:
            continue  # notas al pie del formato, no son temas
        data = {
            "tema": tema,
            "meses": months,
            "modalidad": modalidad,
            "estado": "Programada",
        }
        if section_key == "capacitacion":
            perfil = _text(_merged_value(ws, r, perfil_col)) if perfil_col else ""
            last_perfil = perfil or last_perfil
            data["perfil"] = last_perfil
        rows.append((section_key, data))
    return rows


def import_plan(wb):
    rows = []
    for ws in wb.worksheets:
        name = _norm(ws.title)
        if name.startswith("capacitacion"):
            rows += _plan_sheet(ws, "capacitacion")
        elif name.startswith("concienciacion"):
            rows += _plan_sheet(ws, "concienciacion")
    return rows


# ---------------------------------------------------------------- software autorizado

def import_software(wb):
    ws = wb.worksheets[0]
    header_row, cols = _find_row(ws, lambda vals: any(v.startswith("software externo") for v in vals))
    if not header_row:
        return []
    schema = get("software-autorizado")["sections"][0]
    mapping = {}
    for field in schema["fields"]:
        for alias in field.get("aliases", []):
            col = next((c for k, c in cols.items() if k.startswith(_norm(alias))), None)
            if col:
                mapping[field["key"]] = col
                break
    rows = []
    for r in range(header_row + 1, ws.max_row + 1):
        name = _text(ws.cell(row=r, column=mapping.get("software", 1)).value)
        if not name:
            continue
        data = {key: _text(ws.cell(row=r, column=col).value) for key, col in mapping.items()}
        for field in schema["fields"]:
            data[field["key"]] = clean_value(field, data.get(field["key"], field.get("default", "")))
        rows.append(("software", data))
    return rows


# ---------------------------------------------------------------- obligaciones del OSE

NUMERAL_RE = re.compile(r"^(\d+(?:\.\d+)*|[a-z]\.|[ivx]+\.)\s*", re.IGNORECASE)


def import_ose(wb):
    rows = []
    for ws in wb.worksheets:
        article = re.match(r"\D*(\d+)", ws.title)
        if not article or article.group(1) not in {"5", "6", "7"}:
            continue
        articulo = f"Artículo {article.group(1)}"
        header_row, cols = _find_row(ws, lambda vals: "si" in vals and "no" in vals, limit=5)
        if not header_row:
            continue
        si_col, no_col = cols["si"], cols["no"]
        _, top = _find_row(ws, lambda vals: any(v.startswith("control para") for v in vals), limit=3)
        control_col = next((c for k, c in top.items() if k.startswith("control para")), None)
        resp_col = next((c for k, c in top.items() if k.startswith("responsable")), None)
        for r in range(header_row + 1, ws.max_row + 1):
            texts = [_text(ws.cell(row=r, column=c).value) for c in range(1, si_col)]
            texts = [t for t in texts if t]
            if not texts or texts[-1].upper() == "TOTAL" or texts[0].upper() == "TOTAL":
                continue
            numeral, obligacion = "", " ".join(texts)
            if len(texts) >= 2 and re.fullmatch(r"\d+(\.\d+)+", texts[0]):
                numeral, obligacion = texts[0], " ".join(texts[1:])
            else:
                m = NUMERAL_RE.match(obligacion)
                if m:
                    numeral = m.group(1).rstrip(".")
            si, no = _marked(ws.cell(row=r, column=si_col).value), _marked(ws.cell(row=r, column=no_col).value)
            data = {
                "articulo": articulo,
                "numeral": numeral,
                "obligacion": obligacion,
                "cumple": "Sí" if si else "No" if no else "Por revisar",
                "control": _text(ws.cell(row=r, column=control_col).value) if control_col else "",
                "responsable": _text(ws.cell(row=r, column=resp_col).value) if resp_col else "",
            }
            if not si and not no and re.fullmatch(r"\d+\.\d+", numeral):
                data["tipo"] = "grupo"  # encabezado de un grupo de incisos (5.1, 6.1, 7.7...)
            rows.append(("requisitos", data))
    return rows


def import_system_format(slug, wb):
    """Lee el Excel que descarga el propio sistema: una hoja por sección y encabezados con los
    nombres de los campos. Así cualquier registro exportado se puede editar en Excel y volver a subir."""
    from .schemas import MONTH_LETTERS

    schema = get(slug)
    rows = []
    for section in schema["sections"]:
        names = {_norm(section.get("sheet", "")), _norm(section["label"]), _norm(section["label"][:31])} - {""}
        sheet = next((ws for ws in wb.worksheets if _norm(ws.title) in names), None)
        if sheet is None and len(schema["sections"]) == 1 and len(wb.worksheets) == 1:
            sheet = wb.worksheets[0]
        if sheet is None:
            continue
        labels = {_norm(f["label"]): f for f in section["fields"]}
        header_row, cols = _find_row(sheet, lambda vals: sum(v in labels for v in vals) >= 2)
        if not header_row:
            continue
        mapping = {labels[k]["key"]: c for k, c in cols.items() if k in labels}
        for r in range(header_row + 1, sheet.max_row + 1):
            data = {}
            for field in section["fields"]:
                col = mapping.get(field["key"])
                raw = _text(sheet.cell(row=r, column=col).value) if col else ""
                if field["type"] == "months":
                    letters = [x.strip().upper() for x in raw.split(",") if x.strip()]
                    # las letras se repiten (M, J, A): se toman en orden del año
                    months, start = [], 0
                    for letter in letters:
                        for i in range(start, 12):
                            if MONTH_LETTERS[i] == letter:
                                months.append(i + 1)
                                start = i + 1
                                break
                    data[field["key"]] = months
                else:
                    data[field["key"]] = clean_value(field, raw or field.get("default", ""))
            required = [f["key"] for f in section["fields"] if f.get("required")]
            if required and not any(data.get(k) for k in required):
                continue
            if slug == "obligaciones-ose" and data.get("cumple") == "Por revisar" and re.fullmatch(r"\d+\.\d+", data.get("numeral", "")):
                data["tipo"] = "grupo"
            rows.append((section["key"], data))
    return rows


def import_by_headers(slug):
    """Lectura por encabezados para registros de una tabla. Según la sección puede leer varias hojas
    (sheets: "all" o lista de nombres), guardar el nombre de la hoja en un campo (sheet_field) y tomar
    el año de la hoja («2025») o de un campo de fecha (year_from)."""
    def run(wb):
        schema = get(slug)
        section = schema["sections"][0]
        wanted = section.get("sheets", "first")
        # Los alias del Excel original y, además, la etiqueta del propio campo (así se reconoce
        # también el Excel que descarga el sistema).
        aliases = {_norm(f["label"]): f for f in section["fields"]}
        aliases.update({_norm(a): f for f in section["fields"] for a in f.get("aliases", [])})
        required = [f["key"] for f in section["fields"] if f.get("required")]
        rows, seen = [], set()
        for ws in wb.worksheets:
            if isinstance(wanted, list) and _norm(ws.title) not in {_norm(n) for n in wanted}:
                continue
            header_row, cols = _find_row(ws, lambda vals: sum(any(v.startswith(a) for a in aliases) for v in vals) >= 2)
            if not header_row:
                continue
            mapping = {}
            for label, col in sorted(cols.items(), key=lambda kv: kv[1]):
                field = next((f for a, f in sorted(aliases.items(), key=lambda kv: -len(kv[0])) if label.startswith(a)), None)
                if field and field["key"] not in mapping:
                    mapping[field["key"]] = col
            sheet_year = int(ws.title.strip()) if re.fullmatch(r"\s*20\d\d\s*", ws.title or "") else None
            for r in range(header_row + 1, ws.max_row + 1):
                data = {}
                for field in section["fields"]:
                    if field["key"] == section.get("sheet_field") and field["key"] not in mapping:
                        data[field["key"]] = ws.title.strip()
                        continue
                    col = mapping.get(field["key"])
                    raw = _text(ws.cell(row=r, column=col).value) if col else ""
                    data[field["key"]] = clean_value(field, raw or field.get("default", ""))
                if required and not any(data.get(k) for k in required):
                    continue
                # Algunas hojas son copia de la del año anterior (medidas correctivas 2026 = 2025):
                # una fila idéntica, fuera del nombre de la hoja, se toma una sola vez.
                fingerprint = tuple(sorted((k, str(v)) for k, v in data.items() if k != section.get("sheet_field")))
                if section.get("year_from", "sheet") == "sheet":
                    fingerprint += (("_hoja", ws.title),)  # filas iguales en años distintos son registros distintos
                if fingerprint in seen:
                    continue
                seen.add(fingerprint)
                if schema["by_year"]:
                    year_from = section.get("year_from", "sheet")
                    year = sheet_year if year_from == "sheet" else None
                    if not year and year_from != "sheet":
                        value = data.get(year_from, "")
                        year = int(value[:4]) if value[:4].isdigit() else None
                        # fechas mal digitadas en el original (p. ej. 2028): se usa el año de la hoja
                        if year and not (2015 <= year <= datetime.date.today().year):
                            year = sheet_year
                        year = year or sheet_year
                    if year:
                        data["_year"] = year
                rows.append((section["key"], data))
            if rows and wanted == "first":
                break
        return rows
    return run


# ---------------------------------------------------------------- cambio de claves

ACCOUNT_TYPES = (("supremo", "Supremo"), (" ad", "Directorio activo"), ("anydesk", "AnyDesk"), ("dispositivo", "Dispositivo"))


def import_password_changes(wb):
    """Solo registra si la celda de cada mes tiene algo. El contenido (a veces la clave escrita)
    nunca se lee como texto ni se guarda."""
    rows = []
    for ws in wb.worksheets:
        if not re.fullmatch(r"\s*20\d\d\s*", ws.title or ""):
            continue
        year = int(ws.title.strip())
        header_row, cols = _find_row(ws, lambda vals: sum(v in MONTH_NAMES for v in vals) >= 6)
        if not header_row:
            continue
        month_cols = {MONTH_NAMES[k]: c for k, c in cols.items() if k in MONTH_NAMES}
        first_month_col = min(month_cols.values())
        for r in range(header_row + 1, ws.max_row + 1):
            labels = [_text(ws.cell(row=r, column=c).value) for c in range(1, first_month_col)]
            label = next((t for t in labels if t), "")
            if not label:
                continue
            months = sorted(m for m, c in month_cols.items() if ws.cell(row=r, column=c).value not in (None, ""))
            lower = f" {label.lower()}"
            tipo = next((name for hint, name in ACCOUNT_TYPES if hint in lower), "Otra")
            cuenta = re.split(r"\s+-\s+", label)[0].strip() if tipo != "Otra" else label
            rows.append(("cuentas", {"cuenta": cuenta, "tipo": tipo, "meses": months, "observacion": "", "_year": year}))
    return rows


MONTH_NAMES = {
    "enero": 1, "febrero": 2, "marzo": 3, "abril": 4, "mayo": 5, "junio": 6, "julio": 7, "agosto": 8,
    "setiembre": 9, "septiembre": 9, "octubre": 10, "noviembre": 11, "diciembre": 12,
}


# ---------------------------------------------------------------- actas en Word

def _docx_fields(doc):
    """Pares «etiqueta → valor» de las tablas del acta (las celdas combinadas se repiten en Word)."""
    fields, cells_seen = {}, []
    for table in doc.tables:
        for row in table.rows:
            uniq = []
            for cell in row.cells:
                text = " ".join(cell.text.split())
                if text and (not uniq or uniq[-1] != text):
                    uniq.append(text)
            cells_seen.append(uniq)
            if len(uniq) >= 2 and not uniq[0].endswith(":"):
                fields.setdefault(_norm(uniq[0]), uniq[1])
            for i in range(len(uniq) - 1):
                if uniq[i].endswith(":"):
                    fields.setdefault(_norm(uniq[i]), uniq[i + 1])
    return fields, cells_seen


def _field(fields, *prefixes):
    for key, value in fields.items():
        if any(key.startswith(_norm(p)) for p in prefixes):
            return value
    return ""


def _year_of(date_iso, path):
    if date_iso[:4].isdigit():
        return int(date_iso[:4])
    return year_from_name(path)


def _acta_borrado(path, doc):
    fields, _ = _docx_fields(doc)
    soporte = _field(fields, "datos sobre los soportes", "soporte")
    if not soporte:
        return None
    fecha = clean_value({"type": "date"}, _field(fields, "fecha de borrado", "fecha"))
    return ("actas", {
        "soporte": soporte,
        "fecha": fecha,
        "metodo": _field(fields, "metodo"),
        "responsable": _field(fields, "persona que realizo", "responsable"),
        "archivo": path.rsplit("/", 1)[-1],
        "_year": _year_of(fecha, path),
    })


def _block_text(element):
    return " ".join("".join(x.text or "" for x in element.iter() if x.tag.endswith("}t")).split())


def _acuerdos(doc):
    """Acuerdos en el orden del acta: párrafos y, si los hay, tablas (p. ej. nuevos OSI con su responsable)."""
    out, capture = [], False
    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            text = _block_text(child)
            if not text:
                continue
            if _norm(text).startswith("acuerdos"):
                capture = True
                continue
            if capture:
                if re.match(r"(?i)^(\d\.\s*)?(enlace|grabaci)", text):
                    break
                out.append(f"• {text}")
        elif tag == "tbl" and capture:
            for i, tr in enumerate(x for x in child.iter() if x.tag.endswith("}tr")):
                cells = [_block_text(tc) for tc in tr if tc.tag.endswith("}tc")]
                cells = [c for c in cells if c]
                if cells and i > 0:  # la primera fila es el encabezado de la tabla
                    out.append("   – " + " · ".join(cells))
    return out


def _acta_reunion(path, doc):
    fields, rows = _docx_fields(doc)
    asunto, numero, asistentes = "", "", []
    in_attendance = False
    for uniq in rows:
        head = uniq[0] if uniq else ""
        if head.upper().startswith("ASUNTO"):
            asunto = head.split(":", 1)[-1].strip()
            numero_cell = next((c for c in uniq[1:] if "REUNI" in c.upper()), "")
            m = re.search(r"No\.?\s*(.+)", numero_cell)
            numero = m.group(1).strip() if m else ""
        if _norm(head).startswith("nombre completo"):
            in_attendance = True
            continue
        if in_attendance:
            # Cada asistente ocupa una fila «nombre | cargo»; una fila de una sola celda (AGENDA...) cierra la lista.
            if len(uniq) >= 2 and not head.endswith(":"):
                asistentes.append(" – ".join(uniq[:2]))
            else:
                in_attendance = False
    acuerdos = _acuerdos(doc)
    objetivo = _field(fields, "objetivo de la reunion")
    if not asunto:
        asunto = objetivo
    if not asunto:
        return None
    fecha = clean_value({"type": "date"}, _field(fields, "fecha"))
    return ("reuniones", {
        "numero": numero,
        "fecha": fecha,
        "asunto": asunto,
        "lugar": _field(fields, "lugar"),
        "asistentes": "\n".join(asistentes),
        "acuerdos": "\n".join(acuerdos),
        "estado": "Pendientes",
        "_year": _year_of(fecha, path),
    })


DOCX_READERS = {"actas-borrado": _acta_borrado, "actas-reunion": _acta_reunion}
MAX_ZIP_MEMBERS = 3000
MAX_MEMBER_BYTES = 40 * 1024 * 1024


def read_docx_files(slug, items):
    from docx import Document

    reader = DOCX_READERS[slug]
    rows = []
    for path, blob in sorted(items, key=lambda x: x[0]):
        try:
            row = reader(path, Document(io.BytesIO(blob)))
        except Exception:
            continue  # acta dañada o que no es un acta
        if row:
            rows.append(row)
    return rows


def read_zip(slug, file_obj):
    if slug not in DOCX_READERS:
        return []
    with zipfile.ZipFile(file_obj) as zf:
        members = [i for i in zf.infolist() if not i.is_dir()][:MAX_ZIP_MEMBERS]
        items = [
            (i.filename, zf.read(i))
            for i in members
            if i.filename.lower().endswith(".docx")
            and not i.filename.rsplit("/", 1)[-1].startswith("~$")
            and not re.search(r"(?i)\bcopia\b", i.filename.rsplit("/", 1)[-1])
            and i.file_size <= MAX_MEMBER_BYTES
        ]
    return read_docx_files(slug, items)


IMPORTERS = {
    "copias-respaldo": import_by_headers("copias-respaldo"),
    "cambio-claves": import_password_changes,
    "actas-borrado": lambda wb: [],
    "actas-reunion": lambda wb: [],
    "plan-capacitacion": import_plan,
    "librerias-externas": import_by_headers("librerias-externas"),
    "creacion-usuarios": import_by_headers("creacion-usuarios"),
    "software-autorizado": import_software,
    "obligaciones-ose": import_ose,
}


def _importer_for(slug):
    if slug in IMPORTERS:
        return IMPORTERS[slug]
    return import_by_headers(slug)  # todo registro de tabla declarado en schemas.EXTRA_REGISTERS


def read(slug, file_obj, name=""):
    """Primero intenta el formato original de SiempreSoft; si no lo reconoce, el formato del sistema.
    Las actas en Word se leen desde un .docx suelto o desde la carpeta comprimida en .zip."""
    lower = (name or getattr(file_obj, "name", "") or "").lower()
    if lower.endswith(".zip"):
        return read_zip(slug, file_obj)
    if lower.endswith(".docx"):
        return read_docx_files(slug, [(name, file_obj.read())]) if slug in DOCX_READERS else []
    wb = load_workbook(file_obj, data_only=True)
    rows = _importer_for(slug)(wb)
    return rows or import_system_format(slug, wb)
