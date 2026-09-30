"""Inventario de activos en Excel con el formato de los registros de SiempreSoft."""

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

from apps.registers.exporters import BODY_FONT, BORDER, HEAD_FILL, HEAD_FONT, WRAP, _title_block

from .models import Asset, AssetClass, AssetMovement, Maintenance


def _sheet(wb, title, sheet_title, headers, rows, widths):
    ws = wb.create_sheet(sheet_title[:31])
    _title_block(ws, title, max(len(headers), 4))
    for c, label in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=c, value=label)
        cell.font, cell.fill, cell.border = HEAD_FONT, HEAD_FILL, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r, row in enumerate(rows, start=4):
        for c, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.font, cell.border, cell.alignment = BODY_FONT, BORDER, WRAP
    for c, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:{get_column_letter(len(headers))}{max(3, len(rows) + 3)}"
    return ws


def _responsible(a):
    if a.custodian_id:
        return a.custodian.get_full_name() or a.custodian.username
    return a.custodian_name or a.owner_role


def build_inventory():
    wb = Workbook()
    wb.remove(wb.active)
    assets = Asset.objects.select_related("custodian").order_by("code")

    _sheet(wb, "INVENTARIO DE EQUIPOS", "Equipos",
           ["Código", "Tipo", "Nombre del equipo", "Modelo / placa", "Responsable", "Ubicación", "Estado", "Último inventario", "Fuente"],
           [[a.code, a.asset_type, a.hostname, a.model, _responsible(a), a.location or a.area, a.get_status_display(),
             (a.extra or {}).get("ultimo_inventario"), a.source] for a in assets.filter(asset_class=AssetClass.EQUIPMENT)],
           [15, 20, 24, 28, 26, 30, 14, 12, 30])
    for klass, title, name in ((AssetClass.INFORMATION, "ACTIVOS DE INFORMACIÓN (PRIMARIOS)", "Activos de información"),
                               (AssetClass.SUPPORT, "ACTIVOS DE SOPORTE", "Activos de soporte")):
        _sheet(wb, title, name, ["Código", "Categoría", "Proceso", "Activo", "Propietario", "C", "I", "D", "Valoración", "Fuente"],
               [[a.code, a.asset_type, a.process, a.name, a.owner_role, a.confidentiality, a.integrity, a.availability, a.valuation, a.source]
                for a in assets.filter(asset_class=klass)],
               [9, 22, 26, 40, 30, 6, 6, 6, 11, 30])
    tech_rows = []
    for a in assets.filter(asset_class=AssetClass.TECHNOLOGY):
        risks = (a.extra or {}).get("riesgos_tecnologicos") or [{}]
        for r in risks:
            tech_rows.append([a.code, a.asset_type, a.process, a.name, a.owner_role, r.get("risk", ""), r.get("consequence", ""),
                              r.get("probability", ""), r.get("level", ""), r.get("controls", "")])
    _sheet(wb, "ACTIVOS TECNOLÓGICOS Y SUS RIESGOS", "Activos tecnológicos",
           ["Código", "Tipo", "Proceso", "Activo", "Propietario", "Riesgo", "Consecuencia", "Probabilidad", "Nivel", "Controles"],
           tech_rows, [10, 16, 22, 30, 26, 40, 12, 12, 10, 40])
    _sheet(wb, "DISPOSITIVOS PERSONALES AUTORIZADOS (BYOD)", "BYOD",
           ["Código", "Dispositivo", "Responsable", "Área", "Estado", "Descripción"],
           [[a.code, a.asset_type, _responsible(a), a.area, a.get_status_display(), (a.extra or {}).get("descripcion", "")]
            for a in assets.filter(asset_class=AssetClass.BYOD)], [26, 22, 26, 22, 12, 40])
    moves = AssetMovement.objects.select_related("asset").order_by("-occurred_at")
    _sheet(wb, "MOVIMIENTOS DE ACTIVOS", "Movimientos",
           ["Fecha", "Código", "Movimiento", "Persona", "Origen", "Destino", "Motivo", "Fuente"],
           [[m.occurred_at.date(), m.asset.code, m.get_movement_type_display(), m.person_name, m.origin, m.destination, m.reason, m.source] for m in moves],
           [12, 15, 22, 24, 22, 26, 40, 30])
    reviews = Maintenance.objects.select_related("asset").order_by("-performed_at")
    _sheet(wb, "REVISIONES Y MANTENIMIENTOS", "Revisiones",
           ["Fecha", "Código", "Tipo", "Persona", "Realizada por", "Resultado", "Fuente"],
           [[m.performed_at.date(), m.asset.code, m.maintenance_type, m.person_name, m.technician, m.result, m.source] for m in reviews],
           [12, 15, 26, 24, 26, 60, 30])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
