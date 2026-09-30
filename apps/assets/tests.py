from io import StringIO

from django.contrib.auth import get_user_model
from django.core.management import call_command
from django.test import TestCase
from django.urls import reverse

from apps.assets.models import Asset, AssetClass, AssetMovement, AssetStatus, Maintenance, MovementType
from apps.registers.models import RegisterEntry


class CargaRealDeSiempresoft(TestCase):
    """Carga el paquete real (apps/assets/data y apps/registers/data) y revisa casos concretos."""

    @classmethod
    def setUpTestData(cls):
        call_command("cargar_datos_siempresoft", stdout=StringIO())
        cls.admin = get_user_model().objects.create_user(username="as-admin", business_code="T-AS", is_superuser=True)

    def setUp(self):
        self.client.force_login(self.admin)

    def test_historial_de_un_equipo_real(self):
        cpu = Asset.objects.get(code="SS1-CPU-012")
        moves = list(cpu.movements.filter(movement_type__in=[MovementType.ASSIGNMENT, MovementType.TRANSFER]).order_by("occurred_at"))
        self.assertEqual(moves[0].person_name, "Jose Torres")
        self.assertEqual(moves[-1].person_name, "Jannyna Pancca")
        self.assertEqual(moves[-1].movement_type, MovementType.TRANSFER)
        self.assertEqual(cpu.custodian.email, "jpancca@siempresoft.com")
        self.assertEqual(cpu.status, AssetStatus.ASSIGNED)

    def test_colaboradores_se_crean_sin_acceso(self):
        user = get_user_model().objects.get(email="jtorres@siempresoft.com")
        self.assertFalse(user.is_active)
        self.assertFalse(user.has_usable_password())
        self.assertIn("Microsoft 365", user.source_document)

    def test_salidas_revisiones_y_clases_de_activo(self):
        self.assertTrue(AssetMovement.objects.filter(movement_type=MovementType.CHECKOUT, asset__code="SS1-LAP-001").exists())
        self.assertTrue(Maintenance.objects.filter(maintenance_type__startswith="Revisión periódica").exists())
        self.assertEqual(Asset.objects.filter(asset_class=AssetClass.INFORMATION).count(), 23)
        self.assertEqual(Asset.objects.filter(asset_class=AssetClass.TECHNOLOGY).count(), 41)
        self.assertEqual(Asset.objects.filter(asset_class=AssetClass.DISPOSED).count() + Asset.objects.filter(
            asset_class=AssetClass.EQUIPMENT, status=AssetStatus.RETIRED).count() >= 1, True)

    def test_la_carga_se_puede_repetir_sin_duplicar(self):
        before = (Asset.objects.count(), AssetMovement.objects.count(), Maintenance.objects.count(), RegisterEntry.objects.count())
        call_command("cargar_datos_siempresoft", stdout=StringIO())
        after = (Asset.objects.count(), AssetMovement.objects.count(), Maintenance.objects.count(), RegisterEntry.objects.count())
        self.assertEqual(before, after)

    def test_registros_reales_cargados(self):
        self.assertEqual(RegisterEntry.objects.filter(register="registro-incidentes").count(), 93)
        self.assertEqual(RegisterEntry.objects.filter(register="medidas-correctivas").count(), 363)

    def test_pantallas_de_activos(self):
        page = self.client.get(reverse("assets:list"))
        self.assertContains(page, "SS1-CPU-012")
        detail = self.client.get(reverse("assets:detail", args=["SS1-CPU-012"]))
        self.assertContains(detail, "Jannyna Pancca")
        self.assertContains(detail, "Trazabilidad")
        self.assertContains(detail, "Inventario 2019")
        info = self.client.get(reverse("assets:list") + "?clase=information")
        self.assertContains(info, "AP-01")

    def test_desde_el_perfil_de_la_persona_se_llega_a_la_ficha_del_equipo(self):
        user = get_user_model().objects.get(email="jpancca@siempresoft.com")
        page = self.client.get(reverse("dashboard:user_profile", args=[user.pk]))
        self.assertContains(page, "SS1-CPU-012")
        self.assertContains(page, reverse("assets:detail", args=["SS1-CPU-012"]))

    def test_inventario_en_excel_con_todas_sus_hojas(self):
        import io

        from openpyxl import load_workbook

        response = self.client.get(reverse("assets:export"))
        self.assertEqual(response.status_code, 200)
        wb = load_workbook(io.BytesIO(response.content))
        for sheet in ("Equipos", "Activos de información", "Activos de soporte", "Activos tecnológicos", "BYOD", "Movimientos", "Revisiones"):
            self.assertIn(sheet, wb.sheetnames)
        codes = [r[0] for r in wb["Equipos"].iter_rows(min_row=4, values_only=True)]
        self.assertIn("SS1-CPU-012", codes)

