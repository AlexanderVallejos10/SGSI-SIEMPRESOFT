import io
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from apps.registers.exporters import BODY_FONT, BORDER, HEAD_FILL, HEAD_FONT, WRAP, _title_block

from .logic import compliance
from .models import FactorGroup, MatrixType
from .selectors import build_dashboard_context

CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)
RELLENO = {"P": PatternFill("solid", fgColor="D9731F"), "S": PatternFill("solid", fgColor="FBE3CC")}
BLANCA = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
CUMPLE = {"SI": PatternFill("solid", fgColor="E2F0D9"), "NO": PatternFill("solid", fgColor="FBE2D5")}


def _tabla(ws, inicio, encabezados, filas, anchos):
    for col, (texto, ancho) in enumerate(zip(encabezados, anchos), start=1):
        celda = ws.cell(row=inicio, column=col, value=texto)
        celda.font, celda.fill, celda.border, celda.alignment = HEAD_FONT, HEAD_FILL, BORDER, CENTRO
        ws.column_dimensions[get_column_letter(col)].width = ancho
    for r, fila in enumerate(filas, start=inicio + 1):
        for c, valor in enumerate(fila, start=1):
            celda = ws.cell(row=r, column=c, value=valor)
            celda.font, celda.border, celda.alignment = BODY_FONT, BORDER, WRAP
    ws.freeze_panes = ws.cell(row=inicio + 1, column=1)
    return inicio + len(filas)


def _numero(valor):
    if valor is None:
        return None
    return float(valor) if isinstance(valor, Decimal) else valor


def _hoja_indicadores(wb, ctx):
    ws = wb.active
    ws.title = "Dashboard SGSI"
    _title_block(ws, "INDICADORES DEL SGSI", 12)
    filas = [[
        f"M{r['obj'].metric_id}", r["obj"].description, r["obj"].pdca_cycle, r["obj"].sgsi_process, r["obj"].method,
        r["obj"].objective, r["obj"].responsible_text, r["obj"].period, r["obj"].indicator,
        _numero(r["effective"]) if r["effective"] is not None else (r["obj"].current_text or ""), r["compliance"], r["obj"].action_plan,
    ] for r in ctx["sgsi_rows"]]
    fin = _tabla(ws, 5, ["N.°", "Indicador", "PHVA", "Proceso del SGSI", "Método", "Objetivo", "Responsable", "Periodo",
                         "Meta", "Valor actual", "¿Cumple?", "Plan de acción"], filas, [7, 44, 8, 22, 34, 30, 22, 14, 12, 12, 10, 34])
    for r in range(6, fin + 1):
        celda = ws.cell(row=r, column=11)
        celda.alignment = CENTRO
        if celda.value in CUMPLE:
            celda.fill = CUMPLE[celda.value]


def _hoja_oesi(wb, ctx):
    ws = wb.create_sheet("Dashboard OESI")
    _title_block(ws, "MEDICIÓN DE OBJETIVOS DE SEGURIDAD", 8)
    filas = [[
        r["obj"].metric_id, r["obj"].description, r["obj"].method, r["obj"].responsible_text, r["obj"].period,
        r["obj"].indicator, _numero(r["obj"].current_value) if r["obj"].current_value is not None else (r["obj"].current_text or ""),
        compliance(r["obj"].indicator, r["obj"].current_value),
    ] for r in ctx["oesi_rows"]]
    fin = _tabla(ws, 5, ["N.°", "Objetivo de seguridad", "Método", "Responsable", "Periodo", "Meta", "Valor actual", "¿Cumple?"],
                 filas, [7, 48, 36, 22, 14, 12, 12, 10])
    for r in range(6, fin + 1):
        celda = ws.cell(row=r, column=8)
        celda.alignment = CENTRO
        if celda.value in CUMPLE:
            celda.fill = CUMPLE[celda.value]


def _hoja_matriz(wb, titulo_hoja, titulo, filas_obj, etiqueta, texto, alineaciones, clave, ctx, puntajes, esperado, obtenido, pct):
    ws = wb.create_sheet(titulo_hoja)
    security = ctx["security"]
    _title_block(ws, titulo, 2 + len(security))
    fila_enc = 5
    for col, valor in enumerate([etiqueta, "Descripción"] + [s.code for s in security], start=1):
        celda = ws.cell(row=fila_enc, column=col, value=valor)
        celda.font, celda.fill, celda.border, celda.alignment = HEAD_FONT, HEAD_FILL, BORDER, CENTRO
    for col, s in enumerate(security, start=3):
        celda = ws.cell(row=fila_enc + 1, column=col, value=s.description)
        celda.font, celda.border, celda.alignment = Font(name="Calibri", size=9), BORDER, CENTRO
        ws.column_dimensions[get_column_letter(col)].width = 16
    ws.cell(row=fila_enc + 1, column=1).border = BORDER
    ws.cell(row=fila_enc + 1, column=2).border = BORDER
    ws.row_dimensions[fila_enc + 1].height = 90
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 56
    mapa = {(getattr(a, clave + "_id"), a.security_objective_id): a.relation for a in alineaciones}
    fila = fila_enc + 2
    for obj in filas_obj:
        for col, valor in ((1, etiqueta_de(obj)), (2, texto(obj))):
            celda = ws.cell(row=fila, column=col, value=valor)
            celda.font, celda.border, celda.alignment = BODY_FONT, BORDER, WRAP
        for col, s in enumerate(security, start=3):
            rel = mapa.get((obj.pk, s.pk), "")
            celda = ws.cell(row=fila, column=col, value=rel)
            celda.border, celda.alignment = BORDER, CENTRO
            if rel in RELLENO:
                celda.fill = RELLENO[rel]
                celda.font = BLANCA if rel == "P" else HEAD_FONT
        fila += 1
    ws.cell(row=fila, column=2, value="Puntaje (P = 3, S = 1)").font = HEAD_FONT
    for col, s in enumerate(security, start=3):
        celda = ws.cell(row=fila, column=col, value=puntajes[s.code])
        celda.font, celda.border, celda.alignment = HEAD_FONT, BORDER, CENTRO
    fila += 2
    for nombre, valor in (("Puntos mínimos esperados", esperado), ("Puntos obtenidos", obtenido), ("Porcentaje de cumplimiento", round(float(pct), 2))):
        ws.cell(row=fila, column=2, value=nombre).font = HEAD_FONT
        ws.cell(row=fila, column=3, value=valor).font = BODY_FONT
        fila += 1
    ws.freeze_panes = ws.cell(row=fila_enc + 2, column=3)


def etiqueta_de(obj):
    return getattr(obj, "code", None) or getattr(obj, "stakeholder", "")


def _hoja_factores(wb, ctx, matriz):
    resumen = ctx["mefi"] if matriz == MatrixType.MEFI else ctx["mefe"]
    ws = wb.create_sheet(matriz)
    _title_block(ws, "MATRIZ DE EVALUACIÓN DE FACTORES INTERNOS" if matriz == MatrixType.MEFI else "MATRIZ DE EVALUACIÓN DE FACTORES EXTERNOS", 4)
    grupos = dict(FactorGroup.choices)
    fila = 5
    for col, (texto, ancho) in enumerate((("Factor", 60), ("Peso", 10), ("Calificación", 13), ("Puntaje", 11)), start=1):
        celda = ws.cell(row=fila, column=col, value=texto)
        celda.font, celda.fill, celda.border, celda.alignment = HEAD_FONT, HEAD_FILL, BORDER, CENTRO
        ws.column_dimensions[get_column_letter(col)].width = ancho
    fila += 1
    por_grupo = {}
    for factor in resumen["rows"]:
        por_grupo.setdefault(factor.group, []).append(factor)
    for grupo, factores in por_grupo.items():
        ws.cell(row=fila, column=1, value=grupos.get(grupo, grupo).upper()).font = HEAD_FONT
        fila += 1
        for f in factores:
            for col, valor in enumerate((f.description, _numero(f.weight), f.classification, _numero(f.score)), start=1):
                celda = ws.cell(row=fila, column=col, value=valor)
                celda.font, celda.border, celda.alignment = BODY_FONT, BORDER, WRAP if col == 1 else CENTRO
            fila += 1
        ws.cell(row=fila, column=3, value="Subtotal").font = HEAD_FONT
        ws.cell(row=fila, column=4, value=_numero(resumen["groups"].get(grupo, Decimal("0")))).font = HEAD_FONT
        fila += 2
    ws.cell(row=fila, column=1, value="TOTAL").font = HEAD_FONT
    ws.cell(row=fila, column=2, value=_numero(resumen["weight_sum"])).font = HEAD_FONT
    ws.cell(row=fila, column=4, value=_numero(resumen["score_sum"])).font = HEAD_FONT


def _hoja_posicion(wb, ctx):
    ws = wb.create_sheet("EFIEFE")
    _title_block(ws, "POSICIÓN ESTRATÉGICA (EFI / EFE)", 3)
    for fila, (nombre, valor) in enumerate((("Total ponderado EFI", ctx["efi"]), ("Total ponderado EFE", ctx["efe"])), start=5):
        ws.cell(row=fila, column=1, value=nombre).font = HEAD_FONT
        ws.cell(row=fila, column=2, value=round(float(valor), 4)).font = BODY_FONT
    ws.column_dimensions["A"].width = 30


def exportar_desde_sistema():
    ctx = build_dashboard_context()
    wb = Workbook()
    _hoja_indicadores(wb, ctx)
    _hoja_oesi(wb, ctx)
    _hoja_matriz(wb, "Matriz OEE vs OSI", "OBJETIVOS ESTRATÉGICOS VS OBJETIVOS DE SEGURIDAD", ctx["strategic"], "OEE",
                 lambda o: o.description, ctx["oee_alignments"], "strategic_objective", ctx, ctx["oee_scores"],
                 ctx["oee_expected"], ctx["oee_obtained"], ctx["oee_pct"])
    _hoja_matriz(wb, "Matriz RyEPI vs OSI", "EXPECTATIVAS DE PARTES INTERESADAS VS OBJETIVOS DE SEGURIDAD", ctx["requirements"],
                 "Parte interesada", lambda o: o.requirement, ctx["req_alignments"], "requirement", ctx, ctx["req_scores"],
                 ctx["req_expected"], ctx["req_obtained"], ctx["req_pct"])
    _hoja_factores(wb, ctx, MatrixType.MEFI)
    _hoja_factores(wb, ctx, MatrixType.MEFE)
    _hoja_posicion(wb, ctx)
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()
