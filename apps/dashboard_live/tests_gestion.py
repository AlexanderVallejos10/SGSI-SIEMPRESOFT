import io

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from .models import DashboardDataset, OeeOsiAlignment, SecurityObjective, StakeholderRequirement, StrategicObjective
from .selectors import build_dashboard_context


class TableroEditable(TestCase):
    def setUp(self):
        self.admin = get_user_model().objects.create_user(username="gt-admin", business_code="T-GT", is_superuser=True)
        self.client.force_login(self.admin)
        self.ds = DashboardDataset.objects.create(code="DS-G", original_name="t.xlsx", file="x.xlsx", checksum_sha256="2" * 64, is_current=True)
        self.osi = [SecurityObjective.objects.create(dataset=self.ds, code=f"OESI{i}", description=f"Seguridad {i}", source_column="DE"[i - 1], sort_order=i) for i in (1, 2)]
        self.oee = StrategicObjective.objects.create(dataset=self.ds, code="OEE1", description="Marketing", source_row=5)
        StakeholderRequirement.objects.create(dataset=self.ds, source_row=5, stakeholder="Clientes", requirement="Valor agregado")
        for s in self.osi:
            OeeOsiAlignment.objects.create(dataset=self.ds, strategic_objective=self.oee, security_objective=s, relation="S", source_cell="X")

    def test_agregar_objetivo_estrategico_crea_sus_celdas(self):
        r = self.client.post(reverse("dashboard_live:crear_elemento", args=["objetivos-estrategicos"]), {"code": "oee2", "description": "Crecer"})
        self.assertEqual(r.status_code, 302)
        nuevo = StrategicObjective.objects.get(dataset=self.ds, code="OEE2")
        self.assertTrue(nuevo.created_in_system)
        self.assertEqual(OeeOsiAlignment.objects.filter(strategic_objective=nuevo).count(), 2)

    def test_agregar_objetivo_de_seguridad_crea_celdas_en_ambas_matrices(self):
        self.client.post(reverse("dashboard_live:crear_elemento", args=["objetivos-seguridad"]), {"code": "OESI3", "description": "Nuevo"})
        nuevo = SecurityObjective.objects.get(dataset=self.ds, code="OESI3")
        self.assertEqual(nuevo.strategic_alignments.count(), 1)
        self.assertEqual(nuevo.requirement_alignments.count(), 1)

    def test_codigo_repetido_no_se_acepta(self):
        r = self.client.post(reverse("dashboard_live:crear_elemento", args=["objetivos-estrategicos"]), {"code": "OEE1", "description": "x"})
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, "Ya existe otro elemento con este código")

    def test_retirar_y_restaurar_sin_borrar(self):
        url = reverse("dashboard_live:vigencia_elemento", args=["objetivos-seguridad", self.osi[1].pk])
        self.client.post(url, {"vigente": "0"})
        ctx = build_dashboard_context()
        self.assertEqual([s.code for s in ctx["security"]], ["OESI1"])
        self.assertTrue(SecurityObjective.objects.filter(pk=self.osi[1].pk).exists())
        self.client.post(url, {"vigente": "1"})
        self.assertEqual(len(build_dashboard_context()["security"]), 2)

    def test_excel_generado_incluye_lo_agregado(self):
        self.client.post(reverse("dashboard_live:crear_elemento", args=["objetivos-estrategicos"]), {"code": "OEE2", "description": "Crecer"})
        r = self.client.get(reverse("dashboard_live:download"))
        wb = load_workbook(io.BytesIO(r.content))
        valores = [c.value for fila in wb["Matriz OEE vs OSI"].iter_rows() for c in fila]
        self.assertIn("OEE2", valores)
        self.assertIn("MEFI", wb.sheetnames)

    def test_sin_permiso_no_puede_crear(self):
        otro = get_user_model().objects.create_user(username="gt-lector", business_code="T-GL")
        self.client.force_login(otro)
        r = self.client.post(reverse("dashboard_live:crear_elemento", args=["objetivos-estrategicos"]), {"code": "OEE9", "description": "x"})
        self.assertEqual(r.status_code, 403)

    def test_pantalla_de_gestion(self):
        r = self.client.get(reverse("dashboard_live:gestionar", args=["objetivos-seguridad"]))
        self.assertContains(r, "OESI2")
        self.assertContains(r, "Nuevo objetivo de seguridad")
