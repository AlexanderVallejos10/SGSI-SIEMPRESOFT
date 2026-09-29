from datetime import date
from io import BytesIO

from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from openpyxl import Workbook

from apps.documents.models import Document
from apps.organization.models import OrganizationalArea, Position, PositionAssignment
from apps.organization.services import assign_user_to_position, close_assignment
from apps.processes.models import ProcessCategory, ProcessNode
from apps.risks.models import Risk
from apps.traceability.access import can_read_document
from apps.traceability.importer import import_workbook
from apps.traceability.models import DocumentLink, Handover, HandoverItem, ImportBatch
from apps.traceability.services import issue_handover, render_docx, risk_level


@override_settings(MEDIA_ROOT="/tmp/sgsi-test-media")
class TraceabilityFlows(TestCase):
    def setUp(self):
        User = get_user_model()
        self.admin = User.objects.create_user(username="admin", business_code="T-ADM", is_superuser=True)
        self.person = User.objects.create_user(
            username="worker", business_code="T-001", first_name="Test", last_name="Worker"
        )
        self.area = OrganizationalArea.objects.create(code="AREA-T", name="Desarrollo")
        self.position = Position.objects.create(code="POS-T", title="Desarrollador", area=self.area)
        self.client.force_login(self.admin)

    def assign(self, position=None, primary=True):
        return assign_user_to_position(
            assignment=PositionAssignment(
                user=self.person,
                position=position or self.position,
                start_date=date.today(),
                is_primary=primary,
            ),
            actor=self.admin,
        )

    def test_primary_secondary_and_close(self):
        assignment = self.assign()
        other = Position.objects.create(code="P2", title="Comité", max_occupants=0)
        self.assign(other, False)
        self.person.refresh_from_db()
        self.assertEqual(self.person.area, "Desarrollo")
        self.area.name = "Ingeniería"
        self.area.save()
        self.person.refresh_from_db()
        self.assertEqual(self.person.area, "Ingeniería")
        close_assignment(assignment=assignment, actor=self.admin)
        self.person.refresh_from_db()
        self.assertEqual(self.person.area, "")

    def test_area_membership_post_updates_worker(self):
        self.assign()
        self.position.area = None
        self.position.save()
        data = {
            "code": self.area.code,
            "name": self.area.name,
            "color": "#dcecf8",
            "sort_order": 100,
            "is_active": "on",
            "positions_to_link": [str(self.position.pk)],
        }
        response = self.client.post(reverse("organization:area_edit", args=[self.area.pk]), data)
        self.assertEqual(response.status_code, 302)
        self.person.refresh_from_db()
        self.assertEqual(self.person.area, "Desarrollo")

    def test_area_cycle_rejected(self):
        other = OrganizationalArea.objects.create(code="A2", name="Child", parent=self.area)
        self.area.parent = other
        with self.assertRaises(ValidationError):
            self.area.full_clean()

    def test_document_permissions_require_verified_grant(self):
        self.assign()
        d = Document.objects.create(code="D1", title="Restricted", classification="restricted")
        link = DocumentLink.objects.create(
            kind="access", document=d, source_title=d.title, position=self.position
        )
        self.assertFalse(can_read_document(self.person, d))
        link.verified = True
        link.save()
        self.assertTrue(can_read_document(self.person, d))
        self.client.force_login(self.person)
        self.assertEqual(self.client.get(reverse("dashboard:document_detail", args=[d.pk])).status_code, 200)
        link.is_active = False
        link.save()
        self.assertEqual(self.client.get(reverse("dashboard:document_detail", args=[d.pk])).status_code, 403)

    def test_handover_snapshot_immutable_and_docx(self):
        self.assign()
        act = Handover.objects.create(
            user=self.person, responsible=self.admin, kind="entry", occurred_on=date.today()
        )
        HandoverItem.objects.create(
            handover=act, category="material", description="Laptop", completed=True, occurred_on=date.today()
        )
        act = issue_handover(act, self.admin)
        self.area.name = "New name"
        self.area.save()
        self.assertEqual(act.snapshot["areas"], "Desarrollo")
        from docx import Document as Word

        self.assertEqual(self.client.get(reverse("traceability:handover_detail", args=[act.pk])).status_code, 200)
        doc = Word(render_docx(act))
        self.assertIn("SIEMPRESOFT", doc.paragraphs[0].text)
        self.assertEqual(
            self.client.get(reverse("traceability:handover_edit", args=[act.pk])).status_code, 403
        )
        with self.assertRaises(ValidationError):
            issue_handover(act, self.admin)

    def test_worker_cannot_access_other_acta(self):
        act = Handover.objects.create(
            user=self.admin, responsible=self.admin, kind="exit", occurred_on=date.today()
        )
        self.client.force_login(self.person)
        self.assertEqual(
            self.client.get(reverse("traceability:handover_detail", args=[act.pk])).status_code, 403
        )

    def test_save_acta_through_form(self):
        data = {
            "kind": "entry",
            "occurred_on": "2026-09-24",
            "responsible": str(self.admin.pk),
            "notes": "test",
            "items-TOTAL_FORMS": "1",
            "items-INITIAL_FORMS": "0",
            "items-MIN_NUM_FORMS": "0",
            "items-MAX_NUM_FORMS": "100",
            "items-0-category": "material",
            "items-0-description": "Laptop",
            "items-0-inventory_code": "EQ-1",
            "items-0-completed": "on",
            "items-0-occurred_on": "2026-09-24",
            "items-0-delivered_by": "IT",
        }
        r = self.client.post(reverse("traceability:handover_new", args=[self.person.pk]), data)
        self.assertEqual(r.status_code, 302)
        act = Handover.objects.get(user=self.person)
        self.assertEqual(act.items.count(), 1)

    def workbook(self):
        book = Workbook()
        s = book.active
        s.title = "Matriz de riesgos"
        for _ in range(4):
            s.append(["header"])
        s.append(
            [
                "R01",
                "Unknown process",
                "Deliberado",
                "Tecnológico",
                "Amenaza",
                "Evento",
                "DES",
                "Escenario",
                "4 - Crítico",
                "3 - Probable",
                None,
                "Control",
                "Unknown role",
            ]
        )
        book.create_sheet("Tratamiento de riesgos")
        f = BytesIO()
        book.save(f)
        return SimpleUploadedFile("risks.xlsx", f.getvalue())

    def test_import_preview_and_duplicate(self):
        result = import_workbook(self.workbook(), "risks", self.admin)
        self.assertEqual(result["rows"], 1)
        self.assertFalse(Risk.objects.exists())
        self.assertFalse(ImportBatch.objects.exists())
        import_workbook(self.workbook(), "risks", self.admin, apply=True)
        batch = ImportBatch.objects.get()
        self.assertTrue(batch.source.name)
        self.assertIn("risks", batch.source.name)
        result = import_workbook(self.workbook(), "risks", self.admin, apply=True)
        self.assertTrue(result["duplicate"])
        self.assertEqual(Risk.objects.count(), 1)
        self.assertEqual(Risk.objects.first().processes.count(), 0)

    def test_link_selected_risks_to_process(self):
        category = ProcessCategory.objects.create(code="OP", name="Operativos", kind="operational")
        process = ProcessNode.objects.create(code="P-1", name="Desarrollo", category=category)
        risk = Risk.objects.create(code="R-1", process="Desarrollo", scenario="Prueba")
        response = self.client.post(
            reverse("traceability:risk_link_many"),
            {"process": str(process.pk), "risks": [str(risk.pk)]},
        )
        self.assertRedirects(
            response,
            reverse("traceability:risks") + f"?process={process.pk}",
            fetch_redirect_response=False,
        )
        self.assertTrue(process.risks.filter(pk=risk.pk).exists())

    def test_risk_matrix_is_lookup_not_product(self):
        self.assertEqual(risk_level(2, 5), "Medio")
        self.assertEqual(risk_level(4, 5), "Muy alto")
        self.assertEqual(risk_level(5, 5), "Pendiente")

    def test_map_has_no_scope_card(self):
        r = self.client.get(reverse("processes:map"))
        self.assertNotContains(r, 'class="process-reference-grid"')
        self.assertContains(r, "data-fit-map")

    def test_generic_user_editor_cannot_assign_groups(self):
        from apps.dashboard.workbench import _editable_fields

        self.assertNotIn("groups", _editable_fields(get_user_model()))
