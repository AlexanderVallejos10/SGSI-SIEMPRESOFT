from io import BytesIO

from django.core.exceptions import ValidationError
from django.db import transaction
from django.utils import timezone

from .models import Handover

LEVELS = (
    ("Muy bajo", "Muy bajo", "Bajo", "Bajo", "Bajo"),
    ("Muy bajo", "Bajo", "Bajo", "Medio", "Medio"),
    ("Bajo", "Bajo", "Medio", "Alto", "Alto"),
    ("Bajo", "Medio", "Alto", "Alto", "Muy alto"),
)


def risk_level(probability, impact):
    if probability in range(1, 5) and impact in range(1, 6):
        return LEVELS[probability - 1][impact - 1]
    return "Pendiente"


@transaction.atomic
def issue_handover(handover, actor):
    handover = Handover.objects.select_for_update().get(pk=handover.pk)
    if handover.status != "draft":
        raise ValidationError("El acta ya fue emitida.")
    items = list(handover.items.all())
    if not items:
        raise ValidationError("Agregue al menos un recurso.")
    for item in items:
        item.full_clean()
    user = handover.user
    assignments = user.organization_assignments.filter(end_date__isnull=True).select_related("position__area")
    handover.snapshot = {
        "name": user.get_full_name() or user.username,
        "business_code": user.business_code,
        "email": user.email,
        "employment_start": str(user.employment_start or ""),
        "positions": ", ".join(a.position.title for a in assignments),
        "areas": ", ".join(sorted({a.position.area.name for a in assignments if a.position.area_id})),
        "responsible": handover.responsible.get_full_name() or handover.responsible.username,
        "issued_at": timezone.now().isoformat(),
        "items": [
            {
                "category": i.get_category_display(),
                "description": i.description,
                "inventory": i.inventory_code,
                "completed": i.completed,
                "date": str(i.occurred_on or ""),
                "responsible": i.delivered_by,
                "notes": i.notes,
                "asset_id": str(i.asset_id or ""),
                "access_id": str(i.access_id or ""),
            }
            for i in items
        ],
    }
    handover.status = "issued"
    handover.updated_by = actor
    handover.save()
    return handover


def render_docx(handover):
    from docx import Document
    from docx.shared import Inches, Pt

    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.65)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    doc.add_heading("SIEMPRESOFT", 0)
    doc.add_heading(
        "ACTA DE "
        + ("ASIGNACIÓN" if handover.kind == "entry" else "DEVOLUCIÓN")
        + " DE MATERIALES Y ACCESOS",
        1,
    )
    data = handover.snapshot
    doc.add_paragraph(
        f"{handover} | Fecha: {handover.occurred_on:%d/%m/%Y} | {handover.get_status_display()}"
    )
    t = doc.add_table(rows=0, cols=2)
    t.style = "Light Shading Accent 1"
    for k, v in [
        ("Trabajador", data.get("name")),
        ("Código", data.get("business_code")),
        ("Área(s)", data.get("areas")),
        ("Cargo(s)", data.get("positions")),
        ("Fecha de ingreso", data.get("employment_start")),
        ("Responsable", data.get("responsible")),
    ]:
        cells = t.add_row().cells
        cells[0].text = k
        cells[1].text = v or "Sin registro"
    for category in ("Material / equipo", "Aplicación / acceso", "Otro"):
        items = [i for i in data.get("items", []) if i["category"] == category]
        if not items:
            continue
        doc.add_heading(category, 2)
        table = doc.add_table(rows=1, cols=5)
        table.style = "Light Shading Accent 1"
        for c, value in zip(
            table.rows[0].cells,
            ["Recurso / inventario", "Sí / No", "Fecha", "Responsable", "Observaciones"],
            strict=True,
        ):
            c.text = value
        from docx.oxml import OxmlElement

        header = OxmlElement("w:tblHeader")
        table.rows[0]._tr.get_or_add_trPr().append(header)
        for i in items:
            for c, value in zip(
                table.add_row().cells,
                [
                    i["description"] + "\n" + i["inventory"],
                    "Sí" if i["completed"] else "No",
                    i["date"],
                    i["responsible"],
                    i["notes"],
                ],
                strict=True,
            ):
                c.text = value
    doc.add_heading("Observaciones", 2)
    doc.add_paragraph(handover.notes or "Sin observaciones.")
    doc.add_paragraph(
        "La emisión registra el contenido del acta; la conformidad se acredita mediante las firmas de entrega y recepción."
    )
    doc.add_paragraph("\nEntrega / recepción: ____________________      Trabajador: ____________________")
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
