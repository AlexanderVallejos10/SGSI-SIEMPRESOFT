import io

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apps.auditlog.models import AuditLog
from apps.risks.models import Risk, RiskAssessment


class ExportarMatrizDeRiesgos(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.admin = get_user_model().objects.create_user(username="exp-admin", business_code="T-EXP", is_superuser=True)
        risk = Risk.objects.create(code="R57", process="Validación OSE", scenario="Prueba de exportación",
                                   event="Rechazo masivo de comprobantes", threat="Cambio normativo")
        RiskAssessment.objects.create(risk=risk, assessed_at=timezone.now(), probability=3, impact=4, inherent_score=12)

    def test_la_matriz_sale_en_el_formato_original_y_con_el_dato_en_su_fila(self):
        self.client.force_login(self.admin)
        response = self.client.get(reverse("traceability:risks_export"))
        self.assertEqual(response.status_code, 200)
        wb = load_workbook(io.BytesIO(response.content))
        self.assertIn("Matriz de riesgos", wb.sheetnames)
        self.assertIn("Tratamiento de riesgos", wb.sheetnames)
        ws = wb["Matriz de riesgos"]
        rows = [r for r in ws.iter_rows(min_row=5, values_only=True) if r[0] == "R57"]
        self.assertEqual(len(rows), 1)
        self.assertIn("Rechazo masivo de comprobantes", rows[0])
        self.assertTrue(AuditLog.objects.filter(action="export_excel", entity="Risk").exists())

    def test_sin_permiso_no_se_exporta(self):
        user = get_user_model().objects.create_user(username="sin-permiso", business_code="T-SP")
        self.client.force_login(user)
        self.assertEqual(self.client.get(reverse("traceability:risks_export")).status_code, 403)
