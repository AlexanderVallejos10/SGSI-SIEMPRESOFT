"""Exporta la matriz de riesgos en el mismo formato en que la entregó el Oficial de Seguridad.

Se parte de la plantilla original (apps/traceability/data/Matriz_Identificacion_Valoracion_Riesgos_SiempreSoft_2026.xlsx):
se conservan sus hojas de guía, datos maestros, mapa de calor, listas desplegables y fórmulas, y solo se
escriben los datos en las columnas que el importador reconoce. El archivo resultante se puede volver a importar."""

import io
import re
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font

from apps.risks.models import Risk

from .importer import RISK_COLUMNS, TREATMENT_COLUMNS, header_map

TEMPLATE = Path(__file__).resolve().parent / "data" / "Matriz_Identificacion_Valoracion_Riesgos_SiempreSoft_2026.xlsx"
CODE_RE = re.compile(r"^R\d+$")


def _risk_values(risk):
    latest = next(iter(risk.assessments.all()), None)
    owner = risk.owner_position.title if risk.owner_position_id else (risk.owner.get_full_name() if risk.owner_id else "")
    values = {f: getattr(risk, f, "") or "" for f in (
        "process", "origin", "category", "threat", "event", "motivation", "affected_asset_text", "scenario",
        "operational_scenario", "finding_origin", "evidence_reference", "existing_controls", "project_name")}
    values.update(code=risk.code, owner=owner, identification_type=risk.get_identification_type_display(),
                  impact=latest.impact if latest else None, probability=latest.probability if latest else None)
    return values


def _treatment_values(t):
    control = f"{t.control.code} {t.control.name}" if t.control_id else ""
    return {
        "risk_code": t.risk.code, "option": t.option, "control": control, "action": t.action,
        "responsible": (t.responsible.get_full_name() or t.responsible.username) if t.responsible_id else "",
        "start": t.start_date or t.source_start or None, "end": t.due_date or t.source_end or None,
        "resources": t.resources, "status": t.status, "residual_impact": t.residual_impact,
        "residual_probability": t.residual_probability, "closed": t.closed_date, "notes": t.notes,
        "third_party_responsibilities": t.third_party_responsibilities, "third_party": t.third_party,
        "contract_reference": t.contract_reference, "avoidance_method": t.avoidance_method,
        "acceptance_justification": t.acceptance_justification,
    }


def _header(ws, spec, key):
    for r in range(1, 11):
        values = [c.value for c in ws[r]]
        found = header_map(values, spec)
        if key in found:
            return r, found
    return None, {}


def _write(ws, row, columns, values, force=("code", "risk_code")):
    for field, index in columns.items():
        if field not in values:
            continue
        cell = ws.cell(row=row, column=index + 1)
        if isinstance(cell.value, str) and cell.value.startswith("=") and field not in force:
            continue  # no se pisan las fórmulas de la plantilla (salvo el ID, que numera por fila y debe ser el real)
        cell.value = values[field]


def _clear(ws, row, columns, keep=()):
    for field, index in columns.items():
        if field in keep:
            continue
        cell = ws.cell(row=row, column=index + 1)
        if not (isinstance(cell.value, str) and cell.value.startswith("=")):
            cell.value = None


def build_risk_matrix():
    risks = list(Risk.objects.select_related("owner", "owner_position").prefetch_related("assessments", "treatments__control", "treatments__responsible")
                 .order_by("code"))
    if TEMPLATE.exists():
        wb = load_workbook(TEMPLATE)
    else:  # sin la plantilla: mismas columnas en un libro simple
        wb = Workbook()
        wb.active.title = "Matriz de riesgos"
        wb.create_sheet("Tratamiento de riesgos")
        for ws, spec in ((wb["Matriz de riesgos"], RISK_COLUMNS), (wb["Tratamiento de riesgos"], TREATMENT_COLUMNS)):
            for i, (_, prefixes, _) in enumerate(spec, start=1):
                ws.cell(row=4, column=i, value=prefixes[0].capitalize()).font = Font(bold=True)

    # ---- Matriz de riesgos
    ws = wb["Matriz de riesgos"]
    head, cols = _header(ws, RISK_COLUMNS, "code")
    rows_by_code, blank_rows = {}, []
    for r in range(head + 1, ws.max_row + 1):
        code = str(ws.cell(row=r, column=cols["code"] + 1).value or "").strip()
        if CODE_RE.match(code):
            rows_by_code[code] = r
    active = {x.code for x in risks if x.is_active}
    for code, r in rows_by_code.items():
        if code not in active:
            _clear(ws, r, cols, keep=("code",))  # fila de la plantilla sin riesgo vigente: queda libre
            blank_rows.append(r)
    next_row = max(rows_by_code.values(), default=head) + 1
    for risk in (x for x in risks if x.is_active):
        row = rows_by_code.get(risk.code)
        if row is None:
            if blank_rows:
                row = blank_rows.pop(0)
            else:
                row, next_row = next_row, next_row + 1
        _write(ws, row, cols, _risk_values(risk))

    # ---- Tratamiento de riesgos
    ws = wb["Tratamiento de riesgos"]
    head, cols = _header(ws, TREATMENT_COLUMNS, "risk_code")
    if head:
        for r in range(head + 1, ws.max_row + 1):
            _clear(ws, r, cols)
        r = head + 1
        for risk in (x for x in risks if x.is_active):
            for t in risk.treatments.all():
                _write(ws, r, cols, _treatment_values(t))
                for index in cols.values():
                    ws.cell(row=r, column=index + 1).alignment = Alignment(wrap_text=True, vertical="top")
                r += 1

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
