"""Arma el paquete de datos reales de SiempreSoft (activos, personas y su historial) a partir de
los archivos del repositorio de OneDrive. Se ejecuta fuera del sistema; el resultado
(apps/assets/data/siempresoft_activos.json) lo carga el comando `cargar_activos_siempresoft`.

Uso:  python scripts/build_seed_siempresoft.py <carpeta con los archivos> <salida.json>

Reglas:
- Solo se usan datos que están en los archivos. Si una relación no es segura (p. ej. una serie de
  Autopilot sin nombre de equipo coincidente), no se crea.
- No se copian IP, claves, hashes de hardware ni el contenido de celdas con claves.
- Cada dato conserva el archivo de origen para que se pueda verificar.
"""

import csv
import datetime as dt
import glob
import io
import json
import os
import re
import sys
import unicodedata
import zipfile
from collections import defaultdict

from openpyxl import load_workbook

CODE_RE = re.compile(r"SS1-([A-Z]{3})-(\d{3})")
TYPES = {
    "CPU": "Computadora (CPU)", "LAP": "Laptop", "MON": "Monitor", "TEC": "Teclado", "MOU": "Mouse",
    "EST": "Estabilizador", "SUP": "Supresor de picos", "ANT": "Antena WiFi", "CAM": "Cámara web",
    "AUR": "Auriculares", "CAR": "Cargador", "ESC": "Escritorio", "MOV": "Celular", "TAB": "Tablet",
    "STB": "Estabilizador", "COO": "Cooler", "ADA": "Adaptador", "IMP": "Impresora", "TIQ": "Ticketera",
    "PIS": "Pizarra", "PRO": "Proyector", "PRY": "Proyector", "MIC": "Micrófono", "TRI": "Trípode",
    "ARO": "Aro de luz", "ECR": "Ecran",
}
PLACEHOLDERS = {"to be filled by o.e.m.", "default string", "system serial number", "system manufacturer",
                "system product name", "", "none", "n/a"}


def norm(value):
    value = unicodedata.normalize("NFKD", str(value or ""))
    value = "".join(c for c in value if not unicodedata.combining(c)).casefold()
    return " ".join(re.sub(r"[^a-z0-9@._-]+", " ", value).split())


def text(value):
    if value is None:
        return ""
    if isinstance(value, (dt.date, dt.datetime)):
        return value.strftime("%Y-%m-%d")
    return " ".join(str(value).replace("_x000D_", " ").split())


def iso_date(value):
    if isinstance(value, (dt.date, dt.datetime)):
        return value.strftime("%Y-%m-%d")
    m = re.match(r"\s*(\d{1,2})[/.-](\d{1,2})[/.-](\d{4})", str(value or ""))
    if m:
        return f"{m.group(3)}-{int(m.group(2)):02d}-{int(m.group(1)):02d}"
    m = re.match(r"\s*(\d{4})-(\d{2})-(\d{2})", str(value or ""))
    return m.group(0).strip() if m else ""


def find(root, pattern):
    hits = sorted(glob.glob(os.path.join(root, "**", pattern), recursive=True))
    return hits[0] if hits else None


# ----------------------------------------------------------------------------- personas

import difflib

NOT_PEOPLE = re.compile(r"\b(contingencia|sala|impresora|consultoria|servidor|todas|areas?)\b")
ROLE_SPLIT = re.compile(r"\s+[-–(]\s*|\s*[-–]\s+")


class People:
    """Directorio de colaboradores armado con todas las fuentes. Une a la misma persona escrita de
    distintas formas («Ana Karim Salazar – Oficial…», «Karim Salazar», «ksalazar@…», «PC-Karim-Salazar»)."""

    def __init__(self):
        self.items = []

    @staticmethod
    def tokens(name):
        return [t for t in norm(name).replace(".", " ").split() if t not in ("ing", "de", "del", "la")]

    def _same(self, person, tokens):
        mine = person["tokens"]
        if len(tokens) < 2 or len(mine) < 2:
            return False
        # mismo apellido (cualquier apellido) y al menos un nombre igual o casi igual (errores de tipeo)
        last_hit = any(t in mine[1:] for t in tokens[1:]) or tokens[-1] in mine
        if not last_hit:
            return False
        firsts = [t for t in tokens if t not in mine[1:]] or tokens[:1]
        return any(difflib.SequenceMatcher(None, a, b).ratio() >= 0.8 for a in firsts for b in mine)

    def find(self, name="", email=""):
        email = (email or "").lower()
        if email:
            for p in self.items:
                if p["email"] == email:
                    return p
            local = email.split("@")[0]
        else:
            local = ""
        for p in self.items:
            if local and p["email"].split("@")[0] == local:
                return p
        toks = self.tokens(name)
        return next((p for p in self.items if self._same(p, toks)), None)

    def add(self, name="", email="", source="", **extra):
        name = " ".join(str(name or "").split())
        email = (email or "").lower()
        if not name and not email:
            return None
        if NOT_PEOPLE.search(norm(name)):
            return None
        person = self.find(name, email)
        if person is None:
            if not email and len(self.tokens(name)) < 2:
                return None  # un solo nombre no identifica a nadie
            person = {"name": name or email.split("@")[0], "email": email, "area": "", "position": "",
                      "start": "", "end": "", "sources": [], "tokens": self.tokens(name)}
            self.items.append(person)
        if email and not person["email"]:
            person["email"] = email
        if name and len(self.tokens(name)) > len(person["tokens"]) and len(self.tokens(name)) <= 4:
            person["name"], person["tokens"] = name, self.tokens(name)
        for key, value in extra.items():
            if value and not person.get(key):
                person[key] = value
        if source and source not in person["sources"]:
            person["sources"].append(source)
        return person

    def resolve(self, raw, add_source=""):
        """Devuelve el identificador de la persona (correo o nombre) o cadena vacía si el texto no es una persona."""
        raw = text(raw)
        if not raw:
            return ""
        low = raw.lower()
        m = re.search(r"([a-z0-9._-]+)@siempresoft\.(com|net)", low)
        if m:
            p = self.find(email=f"{m.group(1)}@siempresoft.com") or self.add(email=f"{m.group(1)}@siempresoft.com", source=add_source)
            return self.key(p)
        host = re.match(r"(?:pc|laptop|desktop)-([a-z]+)-([a-z]+)", low)
        if host:
            p = self.find(f"{host.group(1)} {host.group(2)}")
            return self.key(p) if p else ""
        first = low.split()[0] if low.split() else ""
        if re.fullmatch(r"[a-z]+", first) and (len(low.split()) == 1 or low.split()[1].startswith("-")):
            p = self.find(email=f"{first}@siempresoft.com")  # usuario sin dominio (wbarrantes, «ecruz -qa»)
            return self.key(p) if p else ""
        name, _, role = ROLE_SPLIT.split(raw, 1)[0], None, ROLE_SPLIT.split(raw, 1)[1:] 
        p = self.find(name)
        if p is None and add_source:
            p = self.add(name, source=add_source, position=(role[0].strip(" )") if role else ""))
        return self.key(p) if p else ""

    @staticmethod
    def key(p):
        if not p:
            return ""
        return p["email"] or "nombre:" + "-".join(p["tokens"])

    def merge_initials(self):
        """Une «avillasis@…» con «Angelo Villasis» cuando el usuario del correo es la inicial más un apellido."""
        email_only = [p for p in self.items if p["email"] and len(p["tokens"]) < 2]
        named = [p for p in self.items if len(p["tokens"]) >= 2]
        for e in email_only:
            local = e["email"].split("@")[0]
            match = [n for n in named if any(local == n["tokens"][0][0] + t for t in n["tokens"][1:])]
            if len(match) != 1:
                continue  # sin coincidencia única no se une a nadie
            n = match[0]
            if n["email"] and n["email"] != e["email"]:
                continue
            n["email"] = e["email"]
            for key in ("area", "position", "start", "end"):
                n[key] = n.get(key) or e.get(key, "")
            n["sources"] += [s for s in e["sources"] if s not in n["sources"]]
            self.items.remove(e)

    def export(self):
        self.merge_initials()
        out = []
        for p in sorted(self.items, key=lambda x: x["name"]):
            out.append({k: v for k, v in p.items() if k != "tokens"} | {"key": self.key(p)})
        return out


def load_people(root):
    people = People()
    path = find(root, "*16*Registro*de*permisos*.xlsx")
    if path:
        ws = load_workbook(path, data_only=True)["Microsoft 365"]
        for row in ws.iter_rows(min_row=4, values_only=True):
            name, email = text(row[0]), text(row[1]).lower()
            if not name or "@" not in email:
                continue
            area, _, position = text(row[2]).partition("/")
            people.add(name, email, os.path.basename(path) + " / Microsoft 365",
                       area=area.strip(), position=position.strip(), start=iso_date(row[3]), end=iso_date(row[4]))
    users = find(root, "01 - Creación de usuarios.xlsx")
    if users:
        ws = load_workbook(users, data_only=True).worksheets[0]
        for row in ws.iter_rows(min_row=4, values_only=True):
            if text(row[1]):
                people.add(text(row[1]), text(row[4]).lower(), os.path.basename(users), position=text(row[2]),
                           start=iso_date(row[6]))
    inv = find(root, "Inventario escritorios y PCs.xlsx")
    if inv:
        # correos del inventario, con el nombre sacado del nombre del equipo (PC-Jorge-Contreras)
        wb = load_workbook(inv, data_only=True)
        for ws in sorted(wb.worksheets, key=lambda w: w.title, reverse=True):
            if not re.fullmatch(r"20\d\d", ws.title.strip()):
                continue
            for row in ws.iter_rows(min_row=2, values_only=True):
                cells = [text(v) for v in row]
                email = next((c.lower() for c in cells if re.fullmatch(r"[a-z0-9._-]+@siempresoft\.(com|net)", c.lower())), "")
                host = next((c for c in cells if re.match(r"(?i)(pc|laptop)-[a-z]+-[a-z]+", c)), "")
                if email:
                    m = re.match(r"(?i)(?:pc|laptop)-([a-z]+)-([a-z]+)", host)
                    name = f"{m.group(1)} {m.group(2)}" if m else ""
                    people.add(name, email.replace(".net", ".com"), f"{os.path.basename(inv)} / {ws.title}")
    # Nombres de las declaraciones firmadas por cada colaborador (nombres de archivo reales).
    for zpath in sorted(glob.glob(os.path.join(root, "*.zip"))):
        try:
            names = zipfile.ZipFile(zpath).namelist()
        except zipfile.BadZipFile:
            continue
        for n in names:
            base = os.path.splitext(os.path.basename(n))[0]
            m = re.match(r"\d+\s*-\s*(?:Declaraci[oó]n de aceptaci[oó]n|DJ)\s*-\s*(.+?)(?:\s*-\s*NO VIGENTE)?$", base, re.I)
            if m and "declaraci" in n.lower():
                people.add(m.group(1), source="Declaraciones firmadas (nombre del archivo)")
    return people, path


def person_index(people):
    return people.resolve


# ----------------------------------------------------------------------------- inventario anual

def load_inventory(root, resolve):
    path = find(root, "Inventario escritorios y PCs.xlsx")
    assets, history = {}, defaultdict(list)
    if not path:
        return assets, history, None
    wb = load_workbook(path, data_only=True)
    # del año más reciente al más antiguo: el nombre de equipo y la placa que quedan son los actuales
    for ws in sorted(wb.worksheets, key=lambda w: w.title, reverse=True):
        if not re.fullmatch(r"20\d\d", ws.title.strip()):
            continue
        year = int(ws.title)
        header = [norm(v) for v in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
        col = lambda *names: next((i for i, h in enumerate(header) if any(h.startswith(n) for n in names)), None)
        c_level, c_env, c_user = col("nivel"), col("ambiente"), col("usuario")
        c_host, c_board = col("pc", "hostname"), col("placa")
        level = env = ""
        for row in ws.iter_rows(min_row=2, values_only=True):
            if c_level is not None and text(row[c_level]):
                level = text(row[c_level])
            if c_env is not None and text(row[c_env]):
                env = text(row[c_env]).replace("\n", " ")
            user_raw = text(row[c_user]) if c_user is not None else ""
            host = text(row[c_host]) if c_host is not None else ""
            board = text(row[c_board]) if c_board is not None else ""
            email = resolve(user_raw)
            # modelo de la laptop: columna sin título que sigue a la placa (p. ej. «LENOVO 83GW - PF5W83PX»)
            laptop_model = text(row[c_board + 1]) if c_board is not None and c_board + 1 < len(row) else ""
            for i, cell in enumerate(row):
                if i in (c_user, c_host, c_board) or (c_board is not None and i == c_board + 1):
                    continue
                for prefix, number in CODE_RE.findall(text(cell)):
                    code = f"SS1-{prefix}-{number}"
                    a = assets.setdefault(code, {"code": code, "prefix": prefix, "type": TYPES.get(prefix, "Equipo"),
                                                 "hostname": "", "board": "", "years": [], "seen_year": year})
                    latest_row = a["seen_year"] == year  # nombre y modelo solo del inventario más reciente del código
                    if latest_row and prefix == "CPU" and host and re.match(r"(?i)(pc|desktop)-", host):
                        a["hostname"] = a["hostname"] or host
                        if board and board.lower() not in PLACEHOLDERS:
                            a["board"] = a["board"] or board
                    if latest_row and prefix == "LAP":
                        if host and re.match(r"(?i)laptop-", host):
                            a["hostname"] = a["hostname"] or host
                        if laptop_model and not CODE_RE.search(laptop_model):
                            a["board"] = a["board"] or laptop_model
                    entry = {"year": year, "user": email, "user_text": "" if email else user_raw,
                             "area": env, "level": level}
                    if entry not in history[code]:
                        history[code].append(entry)
    for code, entries in history.items():
        # si un año ya tiene responsable real, las filas sueltas de ese año («anteriormente jtorres») sobran
        with_user = {e["year"] for e in entries if e["user"]}
        entries[:] = [e for e in entries if e["user"] or e["year"] not in with_user]
        entries.sort(key=lambda e: e["year"])
        assets[code]["years"] = sorted({e["year"] for e in entries})
    return assets, history, path


# ----------------------------------------------------------------------------- salidas y retornos

def load_checkouts(root, resolve):
    path = find(root, "*02*Autorizaciones*retirar*activos*.xlsx")
    moves = []
    if not path:
        return moves, None
    ws = load_workbook(path, data_only=True)["Hoja 1"]
    for row in ws.iter_rows(min_row=4, values_only=True):
        label = text(row[0])
        codes = ["SS1-%s-%s" % c for c in CODE_RE.findall(label)]
        if not codes:
            continue
        for code in codes:
            moves.append({
                "code": code, "label": label, "person": text(row[1]), "email": resolve(row[1]),
                "out": iso_date(row[2]), "term": text(row[3]), "back": iso_date(row[4]),
                "authorized_by": text(row[5]), "medium": text(row[6]),
                "source": os.path.basename(path),
            })
    return moves, path


# ----------------------------------------------------------------------------- revisiones periódicas

REVIEW_DATE_RE = re.compile(r"(\d{8})")


def review_date(name):
    m = REVIEW_DATE_RE.search(name)
    if not m:
        return ""
    s = m.group(1)
    if s.startswith("20"):  # 20240102
        return f"{s[:4]}-{s[4:6]}-{s[6:]}"
    return f"{s[4:]}-{s[2:4]}-{s[:2]}"  # 01102021 (ddmmaaaa)


def iter_docx(root):
    for path in sorted(glob.glob(os.path.join(root, "*.zip"))):
        try:
            zf = zipfile.ZipFile(path)
        except zipfile.BadZipFile:
            continue
        for info in zf.infolist():
            name = info.filename
            if name.lower().endswith(".docx") and "revisi" in name.lower() and "copia" not in name.lower():
                yield path, name, zf.read(info)


def load_reviews(root, resolve):
    from docx import Document

    reviews, seen = [], set()
    for zpath, name, blob in iter_docx(root):
        date = review_date(os.path.basename(name))
        if not date or (date, os.path.basename(name)) in seen:
            continue
        seen.add((date, os.path.basename(name)))
        try:
            doc = Document(io.BytesIO(blob))
        except Exception:
            continue
        for table in doc.tables:
            head = [norm(c.text) for c in table.rows[0].cells]
            if not (head and head[0].startswith("usuario") and any(h.startswith("evidencia") for h in head)):
                continue
            for row in table.rows[1:]:
                cells = [" ".join(c.text.split()) for c in row.cells]
                if not cells or not cells[0]:
                    continue
                reviews.append({
                    "date": date, "person": cells[0], "email": resolve(cells[0]),
                    "evidence": cells[1] if len(cells) > 1 else "",
                    "source": os.path.basename(name),
                })
    return reviews


# ----------------------------------------------------------------------------- inventario de software

def load_software(root, resolve):
    snaps = []
    for path in sorted(glob.glob(os.path.join(root, "*.zip"))):
        try:
            zf = zipfile.ZipFile(path)
        except zipfile.BadZipFile:
            continue
        for info in zf.infolist():
            base = os.path.basename(info.filename)
            m = re.match(r"(?:software-inventory-|export-tvm-machine-software-inventory_)([a-z0-9]+)\.csv$", base, re.I)
            if not m:
                continue
            raw = zf.read(info)
            for enc in ("utf-8-sig", "utf-16", "latin-1"):
                try:
                    txt = raw.decode(enc)
                    break
                except UnicodeDecodeError:
                    continue
            lines = txt.splitlines()
            if len(lines) < 2 or not lines[0].startswith("Machine Software Inventory Export"):
                continue
            stamp = dt.datetime.strptime(lines[0].split(",", 1)[1].strip()[:11], "%d %b %Y").strftime("%Y-%m-%d")
            items = []
            for r in csv.DictReader(io.StringIO("\n".join(lines[1:]))):
                items.append({
                    "name": r.get("Name", ""), "vendor": r.get("Vendor", ""), "version": r.get("Installed Version", ""),
                    "weaknesses": int(r.get("Weaknesses") or 0), "exploit": r.get("Has Exploit") == "True",
                    "eos": bool(r.get("EOS software state") or r.get("EOS version state")),
                })
            key = (m.group(1).lower(), stamp)
            if any((s["machine"], s["date"]) == key for s in snaps):
                continue
            snaps.append({"machine": m.group(1).lower(), "email": resolve(m.group(1)), "date": stamp,
                          "items": items, "source": info.filename.rsplit("/", 2)[-2] + "/" + base})
    return snaps


# ----------------------------------------------------------------------------- activos de información, soporte y tecnológicos

def load_risk_assets(root):
    path = find(root, "01 - Proceso de gestión de riesgos_2026.xlsx") or find(root, "*Proceso de gestión de riesgos_202*.xlsx")
    out = {"primary": [], "support": [], "technology": []}
    if not path:
        return out, None
    wb = load_workbook(path, data_only=True)
    for sheet, kind, prefix in (("Activos primarios", "primary", "AP"), ("Activos de soporte", "support", "AS")):
        ws = wb[sheet]
        for row in ws.iter_rows(min_row=5, values_only=True):
            if not isinstance(row[0], (int, float)) or not text(row[3]):
                continue
            out[kind].append({
                "code": f"{prefix}-{int(row[0]):02d}", "category": text(row[1]), "process": text(row[2]),
                "name": text(row[3]), "owner": text(row[4]),
                "c": text(row[5]), "i": text(row[6]), "a": text(row[7]), "value": text(row[8]),
                "source": f"{os.path.basename(path)} / {sheet}",
            })
    ws = wb["Activos tecnológicos"]
    current = None
    for row in ws.iter_rows(min_row=5, values_only=True):
        rid, kind, process, name, owner = (text(v) for v in row[:5])
        if rid and name:
            current = {"code": f"AT-{rid}", "category": kind, "process": process, "name": name, "owner": owner,
                       "risks": [], "source": f"{os.path.basename(path)} / Activos tecnológicos"}
            out["technology"].append(current)
        if current and text(row[5]):
            current["risks"].append({"risk": text(row[5]), "consequence": text(row[6]), "probability": text(row[7]),
                                     "level": text(row[8]), "controls": text(row[9])})
    return out, path


# ----------------------------------------------------------------------------- bajas y BYOD

def load_disposals(root):
    from docx import Document

    rows = []
    for path in sorted(glob.glob(os.path.join(root, "*.zip"))):
        try:
            zf = zipfile.ZipFile(path)
        except zipfile.BadZipFile:
            continue
        for info in zf.infolist():
            if "Actas de Borrado" not in info.filename or not info.filename.lower().endswith(".docx"):
                continue
            doc = Document(io.BytesIO(zf.read(info)))
            fields = {}
            for t in doc.tables:
                for r in t.rows:
                    cells = [" ".join(c.text.split()) for c in r.cells]
                    uniq = [c for i, c in enumerate(cells) if c and (i == 0 or c != cells[i - 1])]
                    if len(uniq) >= 2:
                        fields.setdefault(norm(uniq[0])[:20], uniq[1])
            soporte = next((v for k, v in fields.items() if k.startswith("datos sobre los sop")), "")
            if soporte:
                rows.append({
                    "support": soporte,
                    "date": iso_date(next((v for k, v in fields.items() if k.startswith("fecha")), "")),
                    "method": next((v for k, v in fields.items() if k.startswith("metodo")), ""),
                    "by": next((v for k, v in fields.items() if k.startswith("persona que")), ""),
                    "source": os.path.basename(info.filename),
                })
    return rows


def load_byod(root, resolve):
    path = find(root, "*07*dispositivos*BYOD*.xlsx")
    out = []
    if not path:
        return out
    ws = load_workbook(path, data_only=True).worksheets[0]
    for row in ws.iter_rows(min_row=4, values_only=True):
        worker = text(row[0])
        if not worker:
            continue
        out.append({"person": worker, "email": resolve(worker), "position": text(row[1]), "area": text(row[2]),
                    "device": text(row[3]), "description": text(row[4]), "date": iso_date(row[5]),
                    "note": text(row[6]), "status": text(row[7]), "source": os.path.basename(path)})
    return out


# ----------------------------------------------------------------------------- cargo actual de cada persona

def split_role(raw):
    parts = ROLE_SPLIT.split(text(raw), 1)
    return (parts[0].strip(), parts[1].strip(" )")) if len(parts) == 2 else (parts[0].strip(), "")


def collect_roles(people, reviews, checkouts, registers_path):
    """Cargo más reciente de cada persona, con fecha y fuente. Sirve para ubicarla en el organigrama."""
    found = defaultdict(list)

    def add(raw_name, role, date, source):
        role = re.sub(r"\s+", " ", role or "").strip(" -–")
        if not role or len(role) < 4:
            return
        key = people.key(people.find(raw_name))
        if key:
            found[key].append({"cargo": role, "fecha": date or "", "fuente": source})

    for r in reviews:
        name, role = split_role(r["person"])
        add(name, role, r["date"], r["source"])
    for c in checkouts:
        name, role = split_role(c["person"])
        add(name, role, c["out"], c["source"])
    if registers_path and os.path.exists(registers_path):
        regs = json.load(open(registers_path, encoding="utf-8"))
        for row in regs.get("actas-reunion", {}).get("rows", []):
            for line in row["data"].get("asistentes", "").split("\n"):
                name, role = split_role(line)
                add(name, role, row["data"].get("fecha", ""), "Acta de reunión " + row["data"].get("numero", ""))
        for row in regs.get("byod-personas", {}).get("rows", []):
            d = row["data"]
            add(d.get("trabajador", ""), d.get("cargo", ""), "", "06 - Lista de cargos BYOD")
    for p in people.items:
        if p.get("position"):
            found[people.key(p)].append({"cargo": p["position"], "fecha": p.get("start", ""), "fuente": "Microsoft 365 / creación de usuarios"})
    for p in people.items:
        roles = sorted(found.get(people.key(p), []), key=lambda r: r["fecha"])
        if roles:
            p["cargo_actual"] = roles[-1]


def main(root, out_path, registers_path="apps/registers/data/registros_siempresoft.json"):
    people, _ = load_people(root)
    resolve = lambda raw: people.resolve(raw, add_source="nombre en revisiones, salidas o BYOD")
    equipment, history, inv_path = load_inventory(root, resolve)
    risk_assets, _ = load_risk_assets(root)
    data = {
        "generated": dt.date.today().isoformat(),
        "note": "Datos reales de SiempreSoft tomados de su repositorio de OneDrive. Cada dato indica su archivo de origen.",
        "people": [],
        "equipment": [dict(a, history=history.get(code, []), source=os.path.basename(inv_path or ""))
                      for code, a in sorted(equipment.items())],
        "checkouts": load_checkouts(root, resolve)[0],
        "reviews": load_reviews(root, resolve),
        "software": load_software(root, resolve),
        "information_assets": risk_assets["primary"],
        "support_assets": risk_assets["support"],
        "technology_assets": risk_assets["technology"],
        "disposals": load_disposals(root),
        "byod": load_byod(root, resolve),
    }
    collect_roles(people, data["reviews"], data["checkouts"], registers_path)
    data["people"] = people.export()
    with open(out_path, "w", encoding="utf-8") as fh:
        json.dump(data, fh, ensure_ascii=False, indent=1)
    return data


if __name__ == "__main__":
    result = main(sys.argv[1], sys.argv[2])
    for key, value in result.items():
        if isinstance(value, list):
            print(f"{key}: {len(value)}")
